"""Bounded read-only provenance intake; originals and detailed output stay private.

No historical lifecycle entrypoints, workbook evaluation, or source code execution.
Source-derived evidence: Pannusch/Schmieder, CC-BY-NC-3.0,
Mendeley DOI 10.17632/y2tz67f6ry.1. Software licensing is separate.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import math
import re
from collections import Counter
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / 'docs/analysis/sci_md_5cqa_assay_002'
DATA = ROOT / 'puckworks/data/pannusch2024'
TASK = 'SCI-MD-5CQA-OBSERVATION-PROVENANCE-001'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def source_root():
    value = os.environ.get('PUCKWORKS_EXTERNAL_DATA_ROOT')
    if value:
        return Path(value)
    config = Path(os.environ.get('XDG_CONFIG_HOME', Path.home() / '.config'))
    return Path(json.loads((config / 'puckworks/data_sources.json').read_text())[
        'sources']['PANNUSCH2024_MENDELEY_FULL_REPOSITORY']['path'])


def admit(root, records, bindings):
    """Independent source admission: one failure never discards another source."""
    result = {}
    for source_id, expected in bindings.items():
        matches = [r for r in records if r['source_id'] == source_id]
        item = {'source_id': source_id, 'expected_sha256': expected,
                'state': None, 'identity_status': 'NOT_CHECKED', 'parse_status': 'NOT_ATTEMPTED'}
        result[source_id] = item
        if len(matches) != 1 or matches[0]['sha256'] != expected:
            item.update(state='not applicable', identity_status='BINDING_FAILURE',
                        reason='source register and inherited exact identity disagree or are ambiguous')
            continue
        path = root / matches[0]['source_relpath']
        item['path'] = str(path)
        try:
            observed = sha256(path)
        except FileNotFoundError as exc:
            item.update(state='not located in the explicitly inspected scope', reason=str(exc))
            continue
        except OSError as exc:
            item.update(state='present but inaccessible here', reason=str(exc))
            continue
        item['observed_sha256'] = observed
        if observed != expected:
            item.update(state='present but not admitted', identity_status='IDENTITY_FAILURE',
                        reason='changed bytes; dependent use blocked')
            continue
        item.update(state='present and inspected', identity_status='MATCH',
                    inspection='bytes authenticated; semantic inspection pending')
    return result


def workbook_metadata(path):
    """All sheets, names and explicit references, including hidden/non-numeric sheets."""
    with ZipFile(path) as archive:
        book = ET.fromstring(archive.read('xl/workbook.xml'))
        sheets = [dict(s.attrib) for s in book.findall('s:sheets/s:sheet', NS)]
        for sheet in sheets:
            sheet.setdefault('state', 'visible')
        names = [{**n.attrib, 'text': n.text} for n in book.findall('s:definedNames/s:definedName', NS)]
        relations = {}
        chart_parts = []
        chart_references = {}
        external_parts = {}
        formula_refs = []
        for member in archive.namelist():
            if member.endswith('.rels'):
                relations[member] = [dict(x.attrib) for x in ET.fromstring(archive.read(member))]
            if member.startswith('xl/charts/') and member.endswith('.xml'):
                chart_parts.append(member)
                if re.fullmatch(r'xl/charts/chart\d+\.xml', member):
                    tree = ET.fromstring(archive.read(member))
                    chart_references[member] = {'references': [n.text for n in tree.iter() if n.tag.endswith('}f')],
                                                'text': [n.text for n in tree.iter() if n.tag.endswith('}t')]}
            if member.startswith('xl/externalLinks/') and member.endswith('.xml'):
                external_parts[member] = archive.read(member).decode()
            if member.startswith('xl/worksheets/sheet') and member.endswith('.xml'):
                tree = ET.fromstring(archive.read(member))
                for cell in tree.findall('.//s:c', NS):
                    f = cell.find('s:f', NS)
                    if f is not None and ('!' in (f.text or '') or '[' in (f.text or '')):
                        formula_refs.append({'part': member, 'cell': cell.get('r'),
                                             'formula': f.text, 'attributes': f.attrib})
        return {'sheets': sheets, 'defined_names': names, 'relationships': relations,
                'chart_parts': chart_parts, 'chart_references': chart_references,
                'external_parts': external_parts, 'explicit_formula_references': formula_refs}


def intake(destination):
    destination = Path(destination).resolve()
    if destination == ROOT or ROOT in destination.parents:
        raise ValueError('Detailed evidence must be outside the public checkout')
    destination.mkdir(parents=True, exist_ok=False)
    bindings = json.loads((DOC / 'SOURCE_IDENTITIES.json').read_text())
    with (DATA / 'source_inputs.csv').open() as f:
        records = list(csv.DictReader(f))
    sources = admit(source_root(), records, bindings['source_files'])
    for item in sources.values():
        if item['identity_status'] != 'MATCH':
            continue
        path = Path(item['path'])
        try:
            if path.suffix == '.xlsx':
                item['metadata'] = workbook_metadata(path)
            elif path.suffix == '.m':
                # Text inspection only; preserve source text privately with line locators.
                item['text_lines'] = path.read_text().splitlines()
            elif path.suffix == '.mat':
                # Container metadata only; accepted array reconciliation is reused.
                from scipy.io import whosmat
                item['variables'] = whosmat(path)
            item.update(parse_status='PASS', inspection='metadata inspected')
        except Exception as exc:
            item.update(parse_status='FAIL', parse_reason=type(exc).__name__ + ': ' + str(exc))
    result = {'task': TASK, 'inherited_task_label': bindings['task'], 'sources': sources}
    (destination / 'intake.json').write_text(json.dumps(result, indent=2, sort_keys=True) + '\n')
    return result


FRACTIONS = (1, 2, 3, 5, 7, 10)
SPILLS = {('FIT', 7, 2), ('FIT', 33, 2), ('FIT', 42, 2)}


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def close(a, b):
    return number(a) and number(b) and math.isclose(a, b, rel_tol=1e-10, abs_tol=1e-12)


def dilution(water, coffee):
    if not number(water) or not number(coffee) or water < 0 or coffee <= 0:
        return None
    return (water + coffee) / coffee


def unique_join(records, key):
    if any(v in (None, '', 'UNKNOWN') for v in key):
        return 'UNKNOWN_KEY', None
    matches = [r for r in records if tuple(r['key']) == tuple(key)]
    if not matches:
        return 'UNMATCHED', None
    if len(matches) != 1:
        return 'AMBIGUOUS', None
    return 'MATCH', matches[0]


def support(value, lo, hi, *, analyte, basis, association, validity):
    """Analytical eligibility deliberately has no predictor-support argument."""
    if validity == 'INVALID_SPILL':
        return 'NOT_APPLICABLE_INVALID_SPILL'
    if (analyte != '5CQA' or basis not in ('diluted_mg_per_kg', 'response_mV_min_per_10')
            or not association or not all(number(x) for x in (value, lo, hi)) or lo > hi):
        return 'UNASSESSABLE'
    return 'BELOW' if value < lo else 'ABOVE' if value > hi else 'WITHIN'


def trendline(pairs):
    """Source chart linear trendline: response-on-concentration, free intercept."""
    if len(pairs) < 2 or not all(number(x) and number(y) for x, y in pairs):
        raise ValueError('incomplete chart pairs')
    xbar = math.fsum(x for x, _ in pairs) / len(pairs)
    ybar = math.fsum(y for _, y in pairs) / len(pairs)
    xx = math.fsum((x-xbar)**2 for x, _ in pairs)
    if xx == 0:
        raise ValueError('constant standard concentrations')
    slope = math.fsum((x-xbar)*(y-ybar) for x, y in pairs) / xx
    return slope, ybar - slope*xbar


def calibration(formulas, cached, path, campaign):
    fs, cs = formulas['CGA'], cached['CGA']
    with ZipFile(path) as z:
        tree = ET.fromstring(z.read('xl/charts/chart2.xml'))
    ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    series = tree.findall('.//c:ser', ns)
    if len(series) != 1:
        raise ValueError('ambiguous calibration chart series')
    chart = series[0]
    refs = [n.text for n in chart.findall('.//c:f', ns)]
    if refs != ['CGA!$D$43:$D$55', 'CGA!$E$43:$E$55']:
        raise ValueError('changed original chart references')
    trends = chart.findall('c:trendline', ns)
    if not trends or any(t.find('c:trendlineType', ns).get('val') != 'linear'
                         or t.find('c:intercept', ns) is not None for t in trends):
        raise ValueError('unsupported source trendline convention')
    pairs, blank_rows, partial_rows = [], [], []
    for r in range(43, 56):
        x, y = cs[f'D{r}'].value, cs[f'E{r}'].value
        if x is None and y is None:
            blank_rows.append(r)
        elif number(x) and number(y):
            pairs.append([x, y])
        else:
            partial_rows.append(r)
    if partial_rows:
        raise ValueError('partial chart pair retained: '+str(partial_rows))
    cached_axes = []
    for axis in ('xVal', 'yVal'):
        cached_axes.append({int(pt.get('idx')): float(pt.find('c:v', ns).text)
                            for pt in chart.findall(f'c:{axis}/c:numRef/c:numCache/c:pt', ns)})
    for i, (x, y) in enumerate(pairs):
        if not close(x, cached_axes[0].get(i)) or not close(y, cached_axes[1].get(i)):
            raise ValueError('chart cache differs from referenced cells')
    standards = []
    stock_checks = []
    for stock in (5, 8):
        expected_formula = f'=B{stock}*1000/(C{stock}/1000)'
        if fs[f'F{stock}'].value != expected_formula:
            raise ValueError('changed stock mass basis')
        solute, water = cs[f'B{stock}'].value, cs[f'C{stock}'].value
        computed = solute*1e6/water
        stock_checks.append({'row': stock, 'solute_g': solute, 'water_g': water,
                             'computed_source_mg_per_kg': computed,
                             'cache_match': close(computed, cs[f'F{stock}'].value)})
    rows = list(range(11, 16 if campaign == 'FIT' else 17)) + list(range(25, 30 if campaign == 'FIT' else 31))
    for i, r in enumerate(rows):
        stock = 5 if r < 20 else 8
        f = fs[f'J{r}'].value
        if f != f'=H{r}*$F${stock}/(H{r}+I{r})':
            raise ValueError('changed working standard convention')
        h, water = cs[f'H{r}'].value, cs[f'I{r}'].value
        # Match documented spreadsheet arithmetic; retain missing raw water separately.
        computed = h*cs[f'F{stock}'].value / (h + (0 if water is None else water))
        area = cs[f'K{r}'].value
        standards.append({'row': r, 'stock_row': stock, 'stock_g': h, 'water_g': water,
                          'water_recorded': water is not None, 'nominal_mg_per_kg': cs[f'D{r}'].value,
                          'source_prepared_mg_per_kg': cs[f'J{r}'].value, 'area': area,
                          'preparation_cache_match': close(computed, cs[f'J{r}'].value),
                          'chart_pair_match': close(pairs[i][0], cs[f'J{r}'].value) and close(pairs[i][1], area),
                          'second_response': cs[f'L{r}'].value if campaign == 'PRED' else None})
    slope, intercept = trendline(pairs)
    reported = (24.513, 78.923) if campaign == 'FIT' else (26.383, -88.067)
    return {'standards': standards, 'stock_checks': stock_checks, 'chart_pairs': pairs,
            'blank_chart_rows': blank_rows, 'trendline_display_count': len(trends),
            'slope': slope, 'intercept': intercept,
            'displayed_coefficient_match': round(slope, 3) == reported[0] and round(intercept, 3) == reported[1],
            'zero_chart_point_physical_origin': 'UNRESOLVED' if campaign == 'PRED' else 'NOT_APPLICABLE',
            'nominal_interval_mg_per_kg': [min(s['nominal_mg_per_kg'] for s in standards), max(s['nominal_mg_per_kg'] for s in standards)],
            'prepared_nonzero_interval_mg_per_kg': [min(s['source_prepared_mg_per_kg'] for s in standards), max(s['source_prepared_mg_per_kg'] for s in standards)],
            'chart_concentration_interval_mg_per_kg': [min(x for x, _ in pairs), max(x for x, _ in pairs)],
            'chart_response_interval': [min(y for _, y in pairs), max(y for _, y in pairs)],
            'validated_range': None, 'calibration_sample_association': None,
            'reason': 'Equation provenance alone supplies no explicit analytical batch/sample applicability record'}


def preparation_records(sheet, campaign):
    """Read bounded No./Fr. blocks, preserving duplicate and missing keys."""
    result = []
    for row in sheet:
        for anchor in row:
            if anchor.value != 'No.':
                continue
            r0, col = anchor.row, anchor.column
            stop = next((r for r in range(r0+1, sheet.max_row+1)
                         if sheet.cell(r, col).value == 'No.'), sheet.max_row+1)
            for c in (col+1, col+5):
                sample = sheet.cell(r0, c).value
                if not isinstance(sample, int) or not 1 <= sample <= (45 if campaign == 'FIT' else 24):
                    continue  # other source sample identities are outside the declared frame
                for r in range(r0+1, stop):
                    if sheet.cell(r, col).value != 'Fr.':
                        continue
                    fraction = sheet.cell(r, c).value
                    if fraction not in FRACTIONS:
                        continue
                    headers = [sheet.cell(r+2, c+i).value for i in range(4)]
                    headers_ok = headers[1] == 'IST H20' and headers[3] == 'IST Kaffee'
                    wc, sc = sheet.cell(r+3, c+1), sheet.cell(r+3, c+3)
                    result.append({'key': [campaign, sample, fraction],
                                   'sample_cell': sheet.cell(r0, c).coordinate,
                                   'fraction_cell': sheet.cell(r, c).coordinate,
                                   'water_cell': wc.coordinate, 'coffee_cell': sc.coordinate,
                                   'water_g': wc.value, 'coffee_g': sc.value,
                                   'quantity_headers_ok': headers_ok})
    return result


def observations(formulas, cached, doe, campaign):
    """Only new preparation provenance; no historical assay-mass recalculation."""
    records = preparation_records(doe['HPLCWeightsAlcaloids'], campaign)
    result = []
    for n in range(1, 46 if campaign == 'FIT' else 25):
        condition, rep = (n-1)//3+1, (n-1)%3+1
        sheet = f'{3*condition-2}-{3*condition}'
        fs, cs = formulas[sheet], cached[sheet]
        exp = doe['ExpSheet']
        shot_key_ok = (exp.cell(n+1, 1).value == n and exp.cell(n+1, 3).value == rep
                       and exp.cell(3*condition-1, 2).value == condition)
        for idx, fraction in enumerate(FRACTIONS):
            outrow = 10 + 6*(rep-1) + idx
            formula = fs[f'Y{outrow}'].value
            match = re.match(r'=\(I(\d+)', str(formula))
            if not match:
                raise ValueError('unrecognized accepted assay reference')
            origin = int(match[1])
            raw_id = cs[f'A{origin}'].value
            identity_ok = bool(re.fullmatch(r'(?:\d{8}_)?P'+str(n)+r'\.'+str(fraction), str(raw_id)))
            row = {'key': [campaign, n, fraction], 'sheet': sheet, 'output_cell': f'Y{outrow}',
                   'sample_cell': f'A{origin}', 'dilution_cell': f'C{origin}', 'sample_label': raw_id,
                   'validity': 'INVALID_SPILL' if (campaign, n, fraction) in SPILLS else 'VALID',
                   'role': 'q1' if fraction == 1 else 'q2' if fraction == 2 else 'suffix',
                   'source_key_match': identity_ok and shot_key_ok,
                   'dilution_source': fs[f'C{origin}'].value, 'dilution_cached': cs[f'C{origin}'].value,
                   'shared_preparation_C_T_V_Y': all(f'*$C{origin}*' in str(fs[f'{column}{outrow}'].value)
                                                       for column in ('T', 'V', 'Y'))}
            state, prep = unique_join(records, row['key'])
            row.update(join_state=state, preparation=prep)
            if row['validity'] == 'INVALID_SPILL':
                row['preparation_status'] = 'NOT_APPLICABLE_INVALID_SPILL'
            elif not row['source_key_match'] or state != 'MATCH':
                row['preparation_status'] = 'UNASSESSABLE_IDENTITY'
            elif campaign == 'PRED' and n <= 3:
                row['preparation_status'] = 'REUSED_C01' if row['shared_preparation_C_T_V_Y'] else 'UNASSESSABLE_SHARED_PREPARATION'
            else:
                expected = dilution(prep['water_g'], prep['coffee_g']) if prep['quantity_headers_ok'] else None
                row['independent_dilution'] = expected
                if expected is None or not number(row['dilution_cached']):
                    row['preparation_status'] = 'UNASSESSABLE_QUANTITY'
                else:
                    row['preparation_status'] = 'CONSISTENT' if close(expected, row['dilution_cached']) else 'CONFLICT'
                    if campaign == 'FIT':
                        fm = re.fullmatch(r'=\(([\d.]+)\+([\d.]+)\)/([\d.]+)', str(row['dilution_source']))
                        row['literal_quantities_match'] = bool(fm and float(fm[1]) == prep['water_g']
                                                              and float(fm[2]) == float(fm[3]) == prep['coffee_g'])
                        if not row['literal_quantities_match']:
                            row['preparation_status'] = 'CONFLICT' if fm else 'UNASSESSABLE_FORMULA'
                    if row['preparation_status'] == 'CONFLICT':
                        row['source_value_diagnostic'] = {
                            'label': 'DOCUMENTED_PREPARATION_TO_RETAINED_DILUTION_RATIO_NOT_A_CORRECTION',
                            'multiplier_if_DoE_is_authoritative': expected / row['dilution_cached'],
                            'causal_error_or_model_improvement': None}
            # No assay recalculation for absent calibration applicability.
            row['response_support'] = row['concentration_support'] = (
                'NOT_APPLICABLE_INVALID_SPILL' if row['validity'] == 'INVALID_SPILL' else 'UNASSESSABLE')
            result.append(row)
    return result, records


def reduce_sources(intake_path, destination):
    import openpyxl
    destination = Path(destination).resolve()
    if destination == ROOT or ROOT in destination.parents:
        raise ValueError('Detailed evidence must be outside the public checkout')
    destination.mkdir(parents=True, exist_ok=False)
    receipt = json.loads(Path(intake_path).read_text())
    # Reauthenticate independently, even after an earlier successful intake.
    with (DATA/'source_inputs.csv').open() as f:
        sources = admit(source_root(), list(csv.DictReader(f)),
                        json.loads((DOC/'SOURCE_IDENTITIES.json').read_text())['source_files'])
    result = {'task': TASK, 'sources': sources, 'calibrations': {}, 'observations': [],
              'preparations': {}, 'component_failures': {}, 'annotations': {}}
    for campaign in ('FIT', 'PRED'):
        h = sources['P24-HPLC-'+campaign]
        d = sources['P24-DOE-'+campaign]
        if h['identity_status'] != 'MATCH':
            result['component_failures'][campaign] = 'HPLC unavailable or identity-failed'
            continue
        fs = cs = db = None
        try:
            fs = openpyxl.load_workbook(h['path'], read_only=False, data_only=False)
            cs = openpyxl.load_workbook(h['path'], read_only=True, data_only=True)
            # Preserve finite text/comment evidence for semantic inspection, no numeric census.
            result['annotations'][campaign] = {
                s.title: [{'cell': c.coordinate, 'text': c.value, 'comment': c.comment.text if c.comment else None}
                          for row in s for c in row if c.comment or (isinstance(c.value, str)
                                                                   and not c.value.startswith('='))] for s in fs}
            try:
                result['calibrations'][campaign] = calibration(fs, cs, h['path'], campaign)
            except (ValueError, TypeError, KeyError) as exc:
                result['component_failures'][campaign+'_calibration'] = str(exc)
            if d['identity_status'] == 'MATCH':
                db = openpyxl.load_workbook(d['path'], read_only=False, data_only=True)
                result['annotations'][campaign+'_DoE'] = {
                    s.title: [{'cell': c.coordinate, 'text': c.value, 'comment': c.comment.text if c.comment else None}
                              for row in s for c in row if c.comment or isinstance(c.value, str)] for s in db}
                rows, preps = observations(fs, cs, db, campaign)
                result['observations'].extend(rows)
                result['preparations'][campaign] = preps
            else:
                result['component_failures'][campaign+'_preparation'] = 'DoE unavailable or identity-failed'
        except Exception as exc:
            result['component_failures'][campaign+'_parse'] = type(exc).__name__+': '+str(exc)
        finally:
            for book in (fs, cs, db):
                if book is not None:
                    book.close()
    # Keep source access failures visible; never turn missing campaigns into zero coverage.
    result['aggregate'] = {}
    for campaign in ('FIT', 'PRED'):
        result['aggregate'][campaign] = {}
        for role in ('q1', 'q2', 'suffix'):
            rows = [r for r in result['observations'] if r['key'][0] == campaign and r['role'] == role]
            result['aggregate'][campaign][role] = dict(Counter(r['preparation_status'] for r in rows))
    result['intake_sha256'] = sha256(intake_path)
    result['intake_task_match'] = receipt['task'] == TASK
    result['source_hashes_unchanged_after'] = {k: sha256(v['path']) == v['expected_sha256']
                                             for k, v in sources.items() if v['identity_status'] == 'MATCH'}
    (destination/'reduction.json').write_text(json.dumps(result, sort_keys=True, indent=2, default=str, allow_nan=False)+'\n')
    return result



def public_projection(reduction_path, intake_path, output):
    """Publish only source identities, aggregate counts and bounded dispositions."""
    result = json.loads(Path(reduction_path).read_text())
    metadata = json.loads(Path(intake_path).read_text())
    out = Path(output)
    out.mkdir(parents=True, exist_ok=True)
    flags = ['GOVERNANCE_CLASS=G1', 'NO_GOVERNING_PHYSICS_CHANGE', 'RESEARCH_ONLY',
             'RETROSPECTIVE', 'TARGET_EXPOSED', 'SOURCE_INTERNAL',
             'PHYSICAL_VALIDATION=NOT_ESTABLISHED', 'ANALYTICAL_UNCERTAINTY=NOT_ESTABLISHED']
    coverage = []
    for campaign, shots in (('FIT', 45), ('PRED', 24)):
        for role in ('q1', 'q2', 'suffix'):
            rows = [r for r in result['observations'] if r['key'][0] == campaign and r['role'] == role]
            counts = Counter(r['preparation_status'] for r in rows)
            intended = shots*(4 if role == 'suffix' else 1)
            invalid = 3 if campaign == 'FIT' and role == 'q2' else 0
            valid = intended-invalid
            assessable = sum(counts[k] for k in ('CONSISTENT', 'CONFLICT', 'REUSED_C01'))
            coverage.append({'campaign': campaign, 'role': role, 'physical_shots': shots,
                             'intended': intended, 'source_valid': valid, 'invalid_spill': invalid,
                             'inspected_identity_and_preparation_positions': len(rows),
                             'exact_sample_key_matches': sum(r['source_key_match'] for r in rows),
                             'annotated_invalid_sample_labels': sum(not r['source_key_match'] and r['validity'] == 'INVALID_SPILL' for r in rows),
                             'preparation_assessable': assessable,
                             'preparation_unassessable': valid-assessable,
                             'new_preparation_consistent': counts['CONSISTENT'],
                             'reused_C01_arithmetic': counts['REUSED_C01'],
                             'preparation_flagged_conflict': counts['CONFLICT'],
                             'range_assessable': 0, 'range_unassessable': valid,
                             'range_below': None, 'range_within': None, 'range_above': None,
                             'range_reason': 'No explicit standard-to-sample analytical batch/applicability relationship',
                             'valid_zero_count_reused_authority': 0})
    calibrations = {}
    for campaign, c in result['calibrations'].items():
        fields = ('slope', 'intercept', 'displayed_coefficient_match', 'nominal_interval_mg_per_kg',
                  'prepared_nonzero_interval_mg_per_kg', 'chart_concentration_interval_mg_per_kg',
                  'chart_response_interval', 'validated_range', 'calibration_sample_association', 'reason')
        calibrations[campaign] = {k:c[k] for k in fields}
        calibrations[campaign].update(
            prepared_standard_records=len(c['standards']), chart_pairs=len(c['chart_pairs']),
            recorded_stock_preparations=len(c['stock_checks']),
            blank_water_entries=sum(not x['water_recorded'] for x in c['standards']),
            unassigned_zero_chart_pairs=int(campaign == 'PRED'),
            source_stock_formula_cache_agreement=all(x['cache_match'] for x in c['stock_checks']),
            working_formula_and_chart_agreement=all(x['preparation_cache_match'] and x['chart_pair_match'] for x in c['standards']),
            purity_isomer_certificate=None, independent_injection_replicates=None,
            source_basis='source-labelled mg/kg; stock denominator is water mass, not solute-plus-water mass')
    summary = {'task': TASK, 'labels': flags, 'coverage': coverage, 'calibration': calibrations,
               'execution_coverage': {'named_originals': 'FULLY_INSPECTED_WITH_DECLARED_LEVELS',
                                      'named_originals_inspected': 8,
                                      'finite_scope': 'PARTIALLY_INSPECTED',
                                      'reason': 'Eight originals and internal relevant references inspected; one explicit external Windows original unavailable'},
               'scientific_disposition': ['SPECIFIC_NEW_SOURCE_PREPARATION_INCONSISTENCIES',
                                          'CALIBRATION_REGRESSIONS_REPRODUCED',
                                          'CAUSAL_BIAS_EXPLANATION_UNRESOLVED',
                                          'BATCH_APPLICABILITY_CONTROLS_AND_MATERIAL_RELATIONS_UNASSESSED'],
               'preparation_conflicts': {'FIT_q1': 1, 'FIT_q2': 1, 'FIT_training_target': 1,
                                        'q1_conflict_outside_matched_cohort_due_to_q2_spill': 1,
                                        'PRED_missing_aliquot_quantity': 1,
                                        'source_value_diagnostics': 'PRIVATE: reduction-003/reduction.json observations[].source_value_diagnostic',
                                        'chosen_correction': None, 'effect_on_predictor_bias': None},
               'preserved_populations': {'FIT_shots':45,'FIT_positions':270,'PRED_shots':24,'PRED_positions':144,
                                        'matched_FIT_shots':42,'matched_FIT_suffix_intended':168,
                                        'matched_FIT_qualified_mass_prefixes':165,'PRED_suffix_positions':96},
               'component_failures': result['component_failures'],
               'private_evidence': {'reduction_sha256': sha256(reduction_path),
                                    'metadata_intake_sha256': sha256(intake_path)},
               'rights': 'Pannusch/Schmieder; CC-BY-NC-3.0; Mendeley DOI 10.17632/y2tz67f6ry.1; software licensing separate'}
    # This projection describes this completed bounded attempt, not an adaptive classifier.
    if len(result['observations']) != 414 or result['component_failures']:
        raise ValueError('Partial run: retain private results; do not publish complete-attempt template')
    sources = {}
    for sid, item in metadata['sources'].items():
        sources[sid] = {k:item[k] for k in ('state','identity_status','parse_status','expected_sha256','observed_sha256')}
        if 'metadata' in item:
            m=item['metadata']
            sources[sid]['sheets']=[{'name':s['name'],'visibility':s['state']} for s in m['sheets']]
            sources[sid]['defined_names']=m['defined_names']
            sources[sid]['inspection_level']='all-sheet metadata/text; only explicit relevant ranges reduced'
        else:
            sources[sid]['inspection_level']='preprocessing text only' if '-M-PREP-' in sid else 'MAT container metadata; accepted array checks reused'
    reused = [
        'docs/analysis/sci_md_5cqa_assay_002/SOURCE_IDENTITIES.json',
        'docs/analysis/sci_md_5cqa_assay_002/SOURCE_COUNTS.json',
        'docs/analysis/sci_md_5cqa_assay_002/RESULTS.json',
        'docs/analysis/sci_md_5cqa_assay_002/INDEPENDENT_REVIEW.md',
        'docs/analysis/sci_md_5cqa_assay_001/RESULTS.md',
        'docs/analysis/sci_md_5cqa_tds_001/RESULTS.md',
        'docs/analysis/sci_md_5cqa_delivery_001/RESULT.md',
        'docs/analysis/val_case_5cqa_decision_resolution_001/RESULTS.json',
        'docs/research/PANNUSCH_RAW_REPRO_001_REPORT.md',
        'docs/research/ESPRESSO_PROGRAMME_REVIEW_2026-10-01.md',
        'docs/research/espresso_programme_review_20261001/SOURCE_NOTES.md']
    audit={'task': TASK, 'sources': sources, 'inherited_source_task_label': metadata['inherited_task_label'],
           'base_commit': 'e6822df8df4465e61ce4616c7a1c577236cab9b2',
           'base_tree': 'b4492544eac40cd0d6d5f3ff856ed9626fe18918',
           'EWP_read_only_main': '7c025527da7e07982428d4c78a3e79bfa29a76ac',
           'reused_evidence_sha256': {r:sha256(ROOT/r) for r in reused},
           'registers_sha256': {p.name:sha256(p) for p in (DATA/'source_inputs.csv', DATA/'analytical_method_register.csv', DATA/'DATA_DICTIONARY.md')},
           'reference_queue': [
               {'identity':'CGA chart2 and standard input cells in both HPLC books','state':'present and inspected','scope':'upstream regression and preparation arithmetic'},
               {'identity':'HPLC fraction sample rows; DoE No./Fr./IST quantity and ExpSheet keys','state':'present and inspected','scope':'414 positions, no source script execution'},
               {'identity':'all workbook sheet/name/formula/relationship metadata and source text','state':'present and inspected','scope':'finite metadata/text inspection; irrelevant numerical ranges not reduced'},
               {'identity':'FIT HPLC externalLink1 embedded cache','state':'present and inspected','scope':'CGA cache empty; populated external cache concerns experiment 46, outside original frame'},
               {'identity':'FIT HPLC externalLink1 original target','state':'present but inaccessible here','qualification':'known referenced original; existence at historical path not reverified','reason':'author Windows path unavailable; no local identity mapping established; no replacement accepted'},
               {'identity':'other-analyte, telemetry and primed/full-sample numerical relationships','state':'not applicable','scope':'metadata disposed; no new numerical investigation'},
               {'identity':'analytical batch/sample and material lot/roast joins','state':'previously inspected and still unresolved','reason':'no explicit matching key or applicability statement in inspected relevant records'},
               {'identity':'eligible recovery/LOD/LOQ/blank and purity controls','state':'not located in the explicitly inspected scope','reason':'zero chart point is not an identified blank; repeat response column has no recorded repeats'}],
           'queue_policy':'finite named identities and explicit internal references; visited once; unavailable external reference not expanded',
           'private_evidence':summary['private_evidence'],
           'original_bytes_unchanged':all(result['source_hashes_unchanged_after'].values())}
    for name, value in [('RESULTS.json',summary),('SOURCE_AUDIT.json',audit)]:
        (out/name).write_text(json.dumps(value,indent=2,sort_keys=True,allow_nan=False)+'\n')
    return summary

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode', choices=['intake', 'reduce', 'project'])
    parser.add_argument('--output', required=True)
    parser.add_argument('--intake')
    parser.add_argument('--reduction')
    args = parser.parse_args()
    if args.mode == 'intake':
        result = intake(args.output)
        print(json.dumps({k: {field: v[field] for field in ('state', 'identity_status', 'parse_status')}
                          for k, v in result['sources'].items()}, indent=2))
    elif args.mode == 'project':
        if not args.intake or not args.reduction:
            parser.error('--intake and --reduction required for project')
        public_projection(args.reduction, args.intake, args.output)
        print('Aggregate projection written')
    else:
        if not args.intake:
            parser.error('--intake required for reduce')
        result = reduce_sources(args.intake, args.output)
        print(json.dumps({'aggregate': result['aggregate'], 'failures': result['component_failures']}, indent=2))


if __name__ == '__main__':
    main()
