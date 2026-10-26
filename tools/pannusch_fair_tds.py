#!/usr/bin/env python3
"""Bounded prepare/fit/freeze/score command; private evidence must be outside Git."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
from pathlib import Path
import sys
import time

import numpy as np
from scipy.io import loadmat
from scipy.optimize import least_squares

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from puckworks.analysis import pannusch_fair_tds as m

BUNDLE = ROOT/'docs/analysis/sci_h1_pannusch_fair_tds_001'
DATA = ROOT/'puckworks/data/pannusch2024'


def read(path):
    return json.loads(Path(path).read_text())


def write(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False)+'\n')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def rows(path):
    with Path(path).open(newline='') as f:
        return list(csv.DictReader(f))


def csvout(path, values):
    with Path(path).open('w', newline='') as f:
        w = csv.DictWriter(f, list(values[0]), lineterminator='\n')
        w.writeheader()
        w.writerows(values)


def comparison_means(shots, observed, predictions):
    return {condition: {'observed': observed[[i for i,p in enumerate(shots) if p['condition']==condition]].mean(axis=0),
            **{arm: y[[i for i,p in enumerate(shots) if p['condition']==condition]].mean(axis=0)
               for arm,y in predictions.items()}} for condition in m.TEST}


def plot_comparison(means):
    from puckworks.public.schema import Producer
    from puckworks.viz.spec import VizSpec
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    spec = VizSpec(id='sci-h1-pannusch-fair-tds-001', title='Conditional fraction TDS', class_=1,
        producer=Producer('tools.pannusch_fair_tds', 'comparison_means', {}),
        badge='PREDICTED', evidence_strength='reference',
        fidelity_ceiling='Source-internal target-exposed campaign-separated condition means; supplied windows/masses; not physical validation.',
        render_fn='tools.pannusch_fair_tds:plot_comparison', components=[],
        caption='Aggregate research comparison; no registered production component or mechanism rendering.',
        outputs=['comparison.png'])
    assert not spec.validate()
    write(BUNDLE/'FIGURE.json', spec.to_dict())
    fig, axes = plt.subplots(2, 2, figsize=(8, 5.8), sharex=True, sharey=True)
    for ax, condition in zip(axes.flat, m.TEST):
        ax.plot(m.FRACTIONS, means[condition]['observed'], 'ko-', label='Observed mean')
        for arm, style in [('F','s-'), ('E','^-'), ('B','d--')]:
            ax.plot(m.FRACTIONS, means[condition][arm], style, label=arm)
        ax.set_title(condition)
        ax.set_xticks(m.FRACTIONS)
        ax.set_xlabel('Source fraction ID (gaps retained)')
        ax.set_ylabel('TDS (%)')
    axes.flat[0].legend(fontsize=8)
    fig.suptitle('Source-internal, target-exposed conditional TDS; three shots per condition', fontsize=10)
    fig.text(.5, .008, spec.badge_line()+' · campaign-separated · physical validation unestablished', ha='center', fontsize=7)
    fig.tight_layout(rect=(0, .025, 1, 1))
    fig.savefig(BUNDLE/'comparison.png', dpi=160)
    plt.close(fig)


def source_root():
    import os
    config = Path(os.environ.get('XDG_CONFIG_HOME', Path.home()/'.config'))/'puckworks/data_sources.json'
    return Path(os.environ.get('PUCKWORKS_EXTERNAL_DATA_ROOT') or
                read(config)['sources']['PANNUSCH2024_MENDELEY_FULL_REPOSITORY']['path'])


def source_shots(training, chemistry=False):
    filename = 'fit_fraction_replicates.csv' if training else 'prediction_fraction_replicates.csv'
    selected = [r for r in rows(DATA/filename) if r['analyte'] == 'TDS' and
                r['condition_id'] in (m.TRAIN if training else m.TEST)]
    reg = {r['shot_id']: r for r in rows(DATA/'experiment_register.csv')}
    output = []
    for shot in sorted({r['shot_id'] for r in selected}):
        group = sorted([r for r in selected if r['shot_id'] == shot], key=lambda r: int(r['fraction_id']))
        assert tuple(int(r['fraction_id']) for r in group) == m.FRACTIONS
        assert all(r['validity'] == 'VALID' and r['concentration_unit'] == 'percent' for r in group)
        g = reg[shot]
        assert float(g['grind_setting']) == 1.7
        if training:
            T = float(g['nominal_temperature_program_id'][6:-1])
            Q = float(g['nominal_flow_program_id'][6:-4])
            assert all(r['fraction_basis'] == 'MEASURED_MASS_G' for r in group)
        else:
            assert all(r['temperature_start_C'] == r['temperature_end_C'] and
                       r['flow_start_mL_s'] == r['flow_end_mL_s'] for r in group)
            T, Q = float(group[0]['temperature_start_C']), float(group[0]['flow_start_mL_s'])
        windows = [[float(r['fraction_start_s']), float(r['fraction_end_s'])] for r in group]
        p = dict(shot=shot, condition=group[0]['condition_id'], windows=windows,
                 masses_g=[float(r['fraction_liquid_g_or_ml']) for r in group],
                 temperature_C=T, flow_label=Q, start_s=windows[0][0], grind=1.7)
        assert p['start_s'] == 0 and min(p['masses_g']) > 0
        if chemistry:
            key = 'concentration_value' if training else 'measured_concentration'
            p['observed'] = [float(r[key]) for r in group]
        output.append(p)
    assert len(output) == (15 if training else 12)
    return output


def check_sources(private):
    prepared, source = read(private/'PREPARED.json'), read(BUNDLE/'SOURCE.json')
    for name in ('training', 'target_inputs'):
        assert sha(private/f'{name}.json') == prepared[f'{name}_sha256']
    assert sha(BUNDLE/'CONFIG.json') == prepared['config_sha256']
    assert sha(BUNDLE/'SOURCE.json') == prepared['source_sha256']
    assert all(sha(DATA/name) == value for name, value in source['exports'].items())
    paths = {r['source_id']: source_root()/r['source_relpath'] for r in rows(DATA/'source_inputs.csv')}
    assert all(sha(paths[r['source_id']]) == r['sha256'] for r in source['source_records'])


def prepare(private):
    """Inspect originals with the established source MAT/RI/design representations."""
    import openpyxl
    if (private/'PREPARED.json').exists():
        raise ValueError('preparation already exists; preserve original receipt')
    train, target = source_shots(True, True), source_shots(False)
    original_receipts, paths = [], {}
    for r in rows(DATA/'source_inputs.csv'):
        if r['source_id'] not in {f'P24-{kind}-{role}' for kind in ('MAT', 'DOE', 'RI', 'M-PREP')
                                 for role in ('FIT', 'PRED')}:
            continue
        path = source_root()/r['source_relpath']
        assert path.is_file() and sha(path) == r['sha256'], r['source_id']
        paths[r['source_id']] = path
        original_receipts.append(dict(source_id=r['source_id'], sha256=r['sha256'],
                                      inspection='hash verified; content inspected as below'))
    join_errors = []
    for role, shots in [('FIT', train), ('PRED', target)]:
        exp = np.atleast_1d(loadmat(paths[f'P24-MAT-{role}'], squeeze_me=True,
                                   struct_as_record=False)['ExperimentalData'])
        for p in shots:
            run = np.atleast_1d(exp[int(p['condition'][-2:])-1].run)[int(p['shot'][-1])-1]
            times, masses = np.ravel(run.tE), np.ravel(run.mE)
            for j, fid in enumerate(m.FRACTIONS):
                a, b = p['windows'][j]
                join_errors.extend([abs(a-(0 if fid == 1 else times[fid-2])),
                                    abs(b-times[fid-1]), abs(p['masses_g'][j]-masses[fid-1])])
                if role == 'FIT':
                    join_errors.append(abs(p['observed'][j]-np.ravel(run.TdS)[j]))
        # Existing qualified adapter is authoritative. Inspect workbook sheets and
        # source preprocessing formulas without refitting or reading target assays.
        for kind in ('RI', 'DOE'):
            wb = openpyxl.load_workbook(paths[f'P24-{kind}-{role}'], read_only=True, data_only=True)
            sheets = {s.title: [s.max_row, s.max_column] for s in wb}
            write(private/f'{kind}_{role}_structure.json', sheets)
            if role == 'FIT' and kind == 'RI':
                # Numeric training RI source worksheet structure/content, no target chemistry.
                write(private/'RI_FIT_inspection.json', [[str(x) for x in row]
                      for row in list(wb['RI_norm'].values)[:4]])
            wb.close()
        text = paths[f'P24-M-PREP-{role}'].read_text(errors='replace')
        write(private/f'preprocessing_{role}.json', [line for line in text.splitlines()
              if any(key in line for key in ('TdS', 'tds', 'mE', 'tE', 'sampleIdx'))])
    assert max(join_errors) <= 5.1e-9
    write(private/'training.json', train)
    write(private/'target_inputs.json', target)
    manifest = dict(task='SCI-H1-PANNUSCH-FAIR-TDS-001', source_records=original_receipts,
                    training=dict(conditions=list(m.TRAIN), physical_shots=15, valid_windows=90),
                    evaluation=dict(conditions=list(m.TEST), physical_shots=12, valid_windows=72),
                    fractions=list(m.FRACTIONS), max_original_export_rounding_difference=max(join_errors),
                    original_inspection='FIT MAT mass/boundary/TDS and PRED MAT mass/boundary joins; FIT/PRED RI and design workbook sheet dimensions; FIT RI sample cells; FIT/PRED preprocessing formulas. PRED chemistry not accessed by preparation. Accepted exports supply qualified observation values.',
                    exports={n: sha(DATA/n) for n in ('fit_fraction_replicates.csv', 'prediction_fraction_replicates.csv', 'experiment_register.csv', 'source_inputs.csv')},
                    physical_flow_label='UNRESOLVED; common Q=flow_label/1000/980',
                    basis='solution-mass TDS percent; rho_common=980 kg/m3',
                    exposure='SOURCE_INTERNAL / CAMPAIGN_SEPARATED / TARGET_EXPOSED',
                    exclusions='other FIT; PRED temperature and flow ramps; no prefixes condition predictions',
                    rights='Pannusch et al., Mendeley 10.17632/y2tz67f6ry.1, CC-BY-NC-3.0')
    write(BUNDLE/'SOURCE.json', manifest)
    write(private/'PREPARED.json', dict(config_sha256=sha(BUNDLE/'CONFIG.json'),
          source_sha256=sha(BUNDLE/'SOURCE.json'), training_sha256=sha(private/'training.json'),
          target_inputs_sha256=sha(private/'target_inputs.json')))
    print('Prepared 15 training shots/90 windows and 12 evaluation shots/72 windows; no evaluation score.', flush=True)


def encode(parameters, names, config):
    values = []
    for name in names:
        lo, hi = config['bounds'][name]
        value = parameters[name]
        if name in ('K_ref', 'A1', 'A2'):
            lo, hi, value = np.log10([lo, hi, value])
        values.append((value-lo)/(hi-lo))
    return np.array(values)


def decode(x, names, fixed, config):
    p = dict(config['source_parameters'], **fixed)
    for value, name in zip(x, names):
        lo, hi = config['bounds'][name]
        if name in ('K_ref', 'A1', 'A2'):
            lo, hi = np.log10([lo, hi])
            p[name] = float(10**(lo+value*(hi-lo)))
        else:
            p[name] = float(lo+value*(hi-lo))
    p['c_s0'] = 1.
    return p


class BudgetExceeded(Exception):
    pass


def fit(private):
    check_sources(private)
    if (private/'fits.json').exists():
        raise ValueError('fits already exist; no repeated campaign')
    c = read(BUNDLE/'CONFIG.json')
    assert sha(BUNDLE/'CONFIG.json') == read(private/'PREPARED.json')['config_sha256']
    shots = read(private/'training.json')
    observed, weight = np.array([p['observed'] for p in shots]), m.weights(shots)
    start_time = time.monotonic()
    allfits, reductions, diagnostic_calls = {}, {}, 0

    def profiled(arm, p):
        if time.monotonic()-start_time > c['elapsed_model_limit_s']:
            raise RuntimeError('eight-hour model wall limit')
        unit, _ = m.predict(arm, p, shots, cells=c['fit_cells'])
        a = m.amplitude(unit, observed, weight, c['bounds']['c_s0'])
        return unit*a, dict(p, c_s0=a)

    for arm in ('E', 'F'):
        names = ['K_ref', 'gamma'] + (['A1', 'B1', 'A2', 'B2'] if arm == 'F' else [])
        fixed, fits = {}, []
        for index, initial in enumerate(c['starts']):
            if arm == 'F' and index == 2:
                # One local, training-only reduction, including inventory profiling.
                best = min(fits, key=lambda f: f['objective'])
                x = encode(best['parameters'], names, c)
                columns = []
                for j in range(len(names)):
                    xp, xm = x.copy(), x.copy()
                    xp[j], xm[j] = min(1., x[j]+.001), max(0., x[j]-.001)
                    yp, _ = profiled(arm, decode(xp, names, fixed, c))
                    ym, _ = profiled(arm, decode(xm, names, fixed, c))
                    diagnostic_calls += 2
                    columns.append(((yp-ym)*np.sqrt(weight)/(xp[j]-xm[j])).ravel())
                J = np.array(columns).T
                norms = np.linalg.norm(J, axis=0)
                keep = [j for j, n in enumerate(names) if n in ('K_ref', 'gamma') or norms[j] >= .01]
                while len(keep) > 2:
                    _, singular, vectors = np.linalg.svd(J[:, keep], full_matrices=False)
                    if singular[-1]/singular[0] >= 1e-4:
                        break
                    removable = [j for j in range(len(keep)) if names[keep[j]] not in ('K_ref', 'gamma')]
                    if not removable:
                        break
                    keep.pop(max(removable, key=lambda j: abs(vectors[-1, j])))
                fixed = {n: c['source_parameters'][n] for j, n in enumerate(names) if j not in keep}
                reductions = dict(nominal_parameter_count=7, effective_free_count=1+len(keep),
                    fixed=fixed, profiled_amplitude=True, optimizer_dimension=len(keep),
                    derivative_norm_pp=dict(zip(names, norms.tolist())),
                    singular_values=np.linalg.svd(J, compute_uv=False).tolist(),
                    sampled_observables='90 training window TDS values, weighted and amplitude-profiled',
                    diagnostic_model_calls=diagnostic_calls, intervals='No CIs; deterministic local diagnostics, not Sobol indices',
                    limitation='Local numerical rank is not global or physical parameter identification')
                names = [names[j] for j in keep]
                print('Training-only reduction:', reductions, flush=True)
            best, calls, begin = None, 0, time.monotonic()

            def residual(x):
                nonlocal best, calls
                if calls >= c['max_calls_per_start']:
                    raise BudgetExceeded
                calls += 1
                pred, parameters = profiled(arm, decode(x, names, fixed, c))
                r = ((pred-observed)*np.sqrt(weight)).ravel()
                objective = float(r@r)
                if best is None or objective < best['objective']:
                    best = dict(parameters=parameters, objective=objective)
                return r

            status = 'CONVERGED'
            def jacobian(x):
                # Absolute scaled-coordinate differences also work at source B1≈0.
                base = residual(x)
                columns = []
                for j in range(len(x)):
                    xp = x.copy()
                    step = c['difference_step'] if x[j] <= 1-c['difference_step'] else -c['difference_step']
                    xp[j] += step
                    columns.append((residual(xp)-base)/step)
                return np.array(columns).T

            try:
                result = least_squares(residual, np.clip(encode(initial, names, c), 1e-9, 1-1e-9),
                    bounds=(0., 1.), jac=jacobian,
                    ftol=c['optimizer_tolerance'], xtol=c['optimizer_tolerance'],
                    gtol=c['optimizer_tolerance'], max_nfev=c['max_calls_per_start'])
                status = str(result.message)
            except BudgetExceeded:
                status = 'ACTUAL_OBJECTIVE_CALL_CAP'
            receipt = dict(arm=arm, start=index+1, calls=calls, elapsed_s=time.monotonic()-begin,
                           status=status, free_names=['c_s0']+names, fixed=fixed, **best)
            fits.append(receipt)
            write(private/f'{arm}_start_{index+1}.json', receipt)
            print(arm, index+1, 'objective', best['objective'], 'calls', calls, 'status', status, flush=True)
        # Initial full fits remain candidates only if no reduction was required.
        eligible = fits if not fixed else fits[2:]
        selected = min(eligible, key=lambda f: f['objective'])
        admitted = [f for f in eligible if f['objective'] <= selected['objective']*1.01]
        allfits[arm] = dict(selected=selected, starts=fits, admitted_starts=[f['start'] for f in admitted])
    baseline = np.mean([np.mean([p['observed'] for p in shots if p['condition'] == condition], axis=0)
                        for condition in m.TRAIN], axis=0).tolist()
    result = dict(arms=allfits, B=baseline, reduction=reductions,
                  model_elapsed_s=time.monotonic()-start_time, diagnostic_calls=diagnostic_calls)
    write(private/'fits.json', result)
    write(BUNDLE/'PARAMETERS.json', {k: v for k, v in result.items() if k != 'B'})


def frozen_paths():
    return [Path(__file__), ROOT/'puckworks/analysis/pannusch_fair_tds.py',
            ROOT/'puckworks/analysis/pannusch_equilibrium_limit.py',
            *[ROOT/'puckworks/models/pannusch2024'/f for f in
              ('solver.py', 'closures.py', 'temperature_history_fv.py')]]


def freeze(private):
    check_sources(private)
    if (BUNDLE/'FREEZE.json').exists():
        raise ValueError('freeze exists; no replacement')
    c, fits = read(BUNDLE/'CONFIG.json'), read(private/'fits.json')
    parameters = read(BUNDLE/'PARAMETERS.json')
    parameters['parameter_accounting'] = {}
    for arm in ('F', 'E'):
        f = fits['arms'][arm]['selected']
        active = [name for name, x in zip(f['free_names'], encode(f['parameters'], f['free_names'], c))
                  if x < 1e-5 or x > 1-1e-5]
        parameters['parameter_accounting'][arm] = dict(nominal=7 if arm=='F' else 3,
            permitted_after_training_rule=len(f['free_names']), optimizer_dimension=len(f['free_names'])-1,
            active_bound_names=active, interior_parameter_count=len(f['free_names'])-len(active),
            inherited='Source centre-grind geometry, phase capacities, diffusion/water closures and source initial-equilibrium form; common source-based nuisance starts. Independent c_s0/K_ref/gamma fits.',
            interpretation='Local rank and boundary location do not prove global or physical identification')
    parameters['parameter_accounting']['B'] = dict(training_fraction_means=6, optimization_calls=0)
    write(BUNDLE/'PARAMETERS.json', parameters)
    train, target = read(private/'training.json'), read(private/'target_inputs.json')
    begin = time.monotonic()
    numerics, predictions = {}, {}
    p = fits['arms']['F']['selected']['parameters']
    grids = [m.predict('F', p, train, cells=n) for n in c['qualification_cells']]
    a, b, d = [x[0] for x in grids]
    allowance = max(1e-8, float(np.max(2*np.abs(d-b)+np.abs(a-b))))
    halves, diag = m.predict('F', p, train, cells=c['prediction_cells'], halves=True)
    temporal = float(np.max(np.abs(halves-d)))
    coarse = float(np.sqrt(np.sum(m.weights(train)*(a-b)**2)))
    fine = float(np.sqrt(np.sum(m.weights(train)*(b-d)**2)))
    if temporal > c['temporal_limit_pp'] or fine >= coarse:
        raise ValueError('training numerical qualification unresolved')
    from puckworks.analysis import pannusch_equilibrium_limit as eq
    first = train[0]
    scenario = m.Scenario(first['temperature_C'], first['flow_label'], p)
    independent, solution = eq.independent_radau(scenario, first['windows'], cells=40)
    standard, _ = m.predict('F', p, [first], cells=40)
    independent_error = float(np.max(np.abs(100*independent/980-standard[0])))
    if independent_error > 1e-6:
        raise ValueError('independent training coordinate check unresolved')
    numerics = dict(F_allowance_pp=allowance, E_allowance_pp=c['numerical_E_allowance_pp'],
        allowance_scope='Training-derived common engineering allowance; not a continuum certificate or total uncertainty',
        training_spatial_budget_pp=c['training_spatial_allowance_pp'],
        numerical_budget_compliance=allowance <= c['training_spatial_allowance_pp'],
        temporal_max_pp=temporal, coarse_change_weighted_rms_pp=coarse, fine_change_weighted_rms_pp=fine,
        independent_Radau_TDS_max_difference_pp=independent_error,
        training_max_conservation_relative=max(v['mass_balance_relative'] for _, group in grids for v in group),
        training_min_phase=min(v['minimum_phase_kg_m3'] for _, group in grids for v in group))
    # Review-requested extension of the same training-only numerical check to
    # every already-admitted fit. No new optimizer, model or target score.
    fine_training = {}
    for arm in ('E', 'F'):
        fine_training[arm] = []
        for f in fits['arms'][arm]['starts']:
            y, _ = m.predict(arm, f['parameters'], train, cells=c['prediction_cells'])
            objective = float(np.sum(m.weights(train)*(y-np.array([q['observed'] for q in train]))**2))
            fine_training[arm].append(dict(start=f['start'], objective_N200=f['objective'], objective_N800=objective))
            if arm == 'F' and f['start'] in fits['arms']['F']['admitted_starts'] and f['start'] != fits['arms']['F']['selected']['start']:
                alternate = [m.predict('F', f['parameters'], train, cells=n) for n in c['qualification_cells']]
                aa, bb, dd = [x[0] for x in alternate]
                alt_allowance = float(np.max(2*np.abs(dd-bb)+np.abs(aa-bb)))
                alt_halves, _ = m.predict('F', f['parameters'], train, cells=c['prediction_cells'], halves=True)
                if np.max(np.abs(dd-alt_halves)) > c['temporal_limit_pp']:
                    raise ValueError('alternate time qualification unresolved')
                if np.sum(m.weights(train)*(bb-dd)**2) >= np.sum(m.weights(train)*(aa-bb)**2):
                    raise ValueError('alternate spatial qualification unresolved')
                allowance = max(allowance, alt_allowance)
                numerics['training_max_conservation_relative'] = max(numerics['training_max_conservation_relative'],
                    max(v['mass_balance_relative'] for _, group in alternate for v in group))
                numerics['training_min_phase'] = min(numerics['training_min_phase'],
                    min(v['minimum_phase_kg_m3'] for _, group in alternate for v in group))
    numerics.update(F_allowance_pp=allowance,
                    numerical_budget_compliance=allowance <= c['training_spatial_allowance_pp'],
                    training_objective_grid_check=fine_training,
                    qualification_scope='All training-admissible F starts; fine-grid objectives reported for every start; selection unchanged')
    # Frozen before target chemistry: selected and training-admissible alternatives.
    for arm in ('E', 'F'):
        predictions[arm] = {}
        for f in fits['arms'][arm]['starts']:
            if f['start'] not in fits['arms'][arm]['admitted_starts']:
                continue
            y, diagnostics = m.predict(arm, f['parameters'], target, cells=c['prediction_cells'])
            predictions[arm][str(f['start'])] = y.tolist()
            write(private/f'target_{arm}_{f["start"]}_numerics.json', diagnostics)
    predictions['B'] = [fits['B'] for _ in target]
    # Grouped training diagnostics use entire physical shots, no role-changing split.
    training_summary = []
    obs = np.array([q['observed'] for q in train])
    for arm in ('E', 'F', 'B'):
        y = np.tile(fits['B'], (len(train), 1)) if arm == 'B' else m.predict(
            arm, fits['arms'][arm]['selected']['parameters'], train, cells=c['prediction_cells'])[0]
        metrics = [dict(condition=q['condition'], **m.shot_metrics(y[i], obs[i], q['masses_g']))
                   for i, q in enumerate(train)]
        for condition in m.TRAIN:
            training_summary.append(dict(arm=arm, condition=condition,
                **m.aggregate([r for r in metrics if r['condition'] == condition])))
    csvout(BUNDLE/'TRAINING.csv', training_summary)
    write(BUNDLE/'NUMERICS.json', numerics)
    write(private/'predictions.json', predictions)
    write(BUNDLE/'FREEZE.json', dict(stage='PRE_SCORE', new_evaluation_scores=0,
        files={str(p.relative_to(ROOT)): sha(p) for p in frozen_paths()},
        bundle_sha256={f: sha(BUNDLE/f) for f in ('PLAN.md', 'CONFIG.json', 'SOURCE.json', 'PARAMETERS.json')},
        private_sha256={n: sha(private/n) for n in ('PREPARED.json', 'training.json', 'target_inputs.json', 'fits.json', 'predictions.json')},
        numerics_sha256=sha(BUNDLE/'NUMERICS.json'),
        model_elapsed_s=fits['model_elapsed_s']+time.monotonic()-begin,
        observation='100*c_l/980; solution-mass TDS; conditional on each original window and measured mass',
        score='Equal-condition mean of equal-shot mass-weighted RMSE; signed and absolute shot bias',
        bands=[0, .1, .25], threshold_factors=[.75, 1., 1.25],
        sensitivities='12 paired whole-shot and 4 whole-condition deletions; independent numerical allowances; training-admissible start envelopes; all nine threshold combinations',
        practical_adequacy='UNESTABLISHED', independent_review='required before score'))
    print('Frozen predictions; no evaluation scores. Numerical allowance:', allowance, flush=True)


def check_freeze(private):
    check_sources(private)
    f = read(BUNDLE/'FREEZE.json')
    assert all(sha(ROOT/p) == value for p, value in f['files'].items())
    assert all(sha(BUNDLE/p) == value for p, value in f['bundle_sha256'].items())
    assert all(sha(private/p) == value for p, value in f['private_sha256'].items())
    assert sha(BUNDLE/'NUMERICS.json') == f['numerics_sha256']
    review = read(BUNDLE/'REVIEW.json')
    assert review['pre_score'] == 'APPROVED' and review['freeze_sha256'] == sha(BUNDLE/'FREEZE.json')


def score(private):
    check_freeze(private)
    if (private/'SCORED.json').exists():
        raise ValueError('one score only; use retained artifacts')
    shots = source_shots(False, True)
    assert [{k: v for k, v in p.items() if k != 'observed'} for p in shots] == read(private/'target_inputs.json')
    # Verify target chemistry against the original accepted MAT only after freeze.
    sr = next(r for r in rows(DATA/'source_inputs.csv') if r['source_id'] == 'P24-MAT-PRED')
    path = source_root()/sr['source_relpath']
    assert sha(path) == sr['sha256']
    exp = np.atleast_1d(loadmat(path, squeeze_me=True, struct_as_record=False)['ExperimentalData'])
    for p in shots:
        run = np.atleast_1d(exp[int(p['condition'][-2:])-1].run)[int(p['shot'][-1])-1]
        assert np.max(np.abs(np.ravel(run.TdS)-p['observed'])) <= 5.1e-9
    fitted, raw, numerical = read(private/'fits.json'), read(private/'predictions.json'), read(BUNDLE/'NUMERICS.json')
    pred = {arm: np.array(raw[arm][str(fitted['arms'][arm]['selected']['start'])]) for arm in ('F', 'E')}
    pred['B'] = np.array(raw['B'])
    obs = np.array([p['observed'] for p in shots])
    metrics = {arm: [dict(shot=p['shot'], condition=p['condition'],
                        **m.shot_metrics(y[i], obs[i], p['masses_g'])) for i, p in enumerate(shots)]
               for arm, y in pred.items()}
    allowances = dict(F=numerical['F_allowance_pp'], E=numerical['E_allowance_pp'], B=0.)
    score_rows = []
    for arm, records in metrics.items():
        for condition in ['ALL']+list(m.TEST):
            r = records if condition == 'ALL' else [x for x in records if x['condition'] == condition]
            score_rows.append(dict(arm=arm, condition=condition, shots=len(r), windows=6*len(r),
                                   **m.aggregate(r), numerical_allowance_pp=allowances[arm]))
    csvout(BUNDLE/'SCORES.csv', score_rows)
    cases = [('all', None)] + [('shot', p['shot']) for p in shots] + [('condition', c) for c in m.TEST]
    deletion = []
    for kind, omit in cases:
        summary = {arm: m.aggregate([r for r in records if kind == 'all' or r[kind] != omit])
                   for arm, records in metrics.items()}
        for candidate, reference in [('F', 'B'), ('E', 'B'), ('F', 'E')]:
            r, q = summary[candidate]['R'], summary[reference]['R']
            ec, er = allowances[candidate], allowances[reference]
            deletion.append(dict(kind=kind, omitted=omit, comparison=f'{candidate}/{reference}',
                gain_pp=q-r, g=1-r/q, g_low=1-(r+ec)/(q-er), g_high=1-max(0,r-ec)/(q+er)))
    write(private/'deletions.json', deletion)
    write(private/'shot_metrics.json', metrics)
    write(private/'observed.json', shots)
    gains, thresholds = [], []
    for comparison in ('F/B', 'E/B', 'F/E'):
        records = [r for r in deletion if r['comparison'] == comparison]
        base = records[0]
        gains.append(dict(comparison=comparison, gain_pp=base['gain_pp'], g=base['g'],
            band=m.band(base['g']), numerical_g_low=base['g_low'], numerical_g_high=base['g_high'],
            shot_delete_g_min=min(r['g'] for r in records if r['kind']=='shot'),
            shot_delete_g_max=max(r['g'] for r in records if r['kind']=='shot'),
            condition_delete_g_min=min(r['g'] for r in records if r['kind']=='condition'),
            condition_delete_g_max=max(r['g'] for r in records if r['kind']=='condition'),
            joint_delete_numerical_g_min=min(r['g_low'] for r in records),
            joint_delete_numerical_g_max=max(r['g_high'] for r in records)))
        for low in (.75, 1., 1.25):
            for high in (.75, 1., 1.25):
                thresholds.append(dict(comparison=comparison, low_factor=low, high_factor=high,
                    baseline_band=m.band(base['g'], .1*low, .25*high),
                    deletion_bands=';'.join(sorted({m.band(r['g'], .1*low, .25*high) for r in records})),
                    deletion_numerical_bands=';'.join(sorted({b for r in records
                        for b in m.interval_bands(r['g_low'], r['g_high'], .1*low, .25*high)}))))
    csvout(BUNDLE/'GAINS.csv', gains)
    csvout(BUNDLE/'THRESHOLDS.csv', thresholds)
    optimization = {}
    for candidate, reference in [('F', 'B'), ('E', 'B'), ('F', 'E')]:
        ranges = {}
        for arm in (candidate, reference):
            alternatives = {'B': raw['B']} if arm == 'B' else raw[arm]
            ranges[arm] = [m.aggregate([dict(condition=p['condition'],
                **m.shot_metrics(y[i], obs[i], p['masses_g'])) for i,p in enumerate(shots)])['R']
                for y in alternatives.values()]
        optimization[f'{candidate}/{reference}'] = [min(1-a/b for a in ranges[candidate] for b in ranges[reference]),
                                                    max(1-a/b for a in ranges[candidate] for b in ranges[reference])]
    spreads = [float(np.std(obs[[i for i,p in enumerate(shots) if p['condition']==c], j], ddof=1))
               for c in m.TEST for j in range(6)]
    result = dict(coverage=dict(expected_windows_per_arm=72, predicted_windows_per_arm=72,
                   physical_shots=12, conditions=4, failures=[]),
        equal_condition={arm: m.aggregate(r) for arm,r in metrics.items()},
        equal_shot={arm: {k: float(np.mean([q[k] for q in r])) for k in ('R','B','A')} for arm,r in metrics.items()},
        gains=gains, optimization_g_ranges=optimization,
        observed_between_shot_SD_pp=dict(min=min(spreads), max=max(spreads), groups=24,
            meaning='Descriptive at differing actual windows; not assay repeatability or CIs'),
        between_shot_metric_SD_pp={arm: {k: float(np.std([q[k] for q in r], ddof=1))
             for k in ('R', 'B', 'A')} for arm,r in metrics.items()},
        numerical_compliance=numerical['numerical_budget_compliance'],
        practical_adequacy='UNESTABLISHED', physical_validation='NOT_ESTABLISHED',
        measurement_training_input_uncertainty='Unquantified, not zero',
        model_elapsed_s=read(BUNDLE/'FREEZE.json')['model_elapsed_s'])
    write(BUNDLE/'RESULTS.json', result)
    # Aggregate figure: condition means only, no replicate traces or raw exports.
    plot_comparison(comparison_means(shots, obs, pred))
    write(private/'SCORED.json', dict(freeze_sha256=sha(BUNDLE/'FREEZE.json'), result_sha256=sha(BUNDLE/'RESULTS.json')))
    print(json.dumps(result, indent=2), flush=True)


def main():
    global BUNDLE
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('stage', choices=['prepare', 'fit', 'freeze', 'score'])
    parser.add_argument('--private', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=BUNDLE,
                        help='Fresh public-output directory for reproduction; source bundle is preserved')
    args = parser.parse_args()
    if args.private.resolve().is_relative_to(ROOT):
        parser.error('private evidence must be outside the repository')
    if args.output.resolve() != BUNDLE.resolve():
        import shutil
        args.output.mkdir(parents=True, exist_ok=True)
        for name in ('PLAN.md', 'CONFIG.json'):
            if not (args.output/name).exists():
                shutil.copy2(BUNDLE/name, args.output/name)
        BUNDLE = args.output
    args.private.mkdir(parents=True, exist_ok=True)
    globals()[args.stage](args.private)


if __name__ == '__main__':
    main()
