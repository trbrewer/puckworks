"""Task-local trigonelline projection and frozen research comparison.

Source originals and row artifacts stay private. Preparation exposes only FIT
chemistry; PRED chemistry crosses the boundary only in the approved scorer.
"""
from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
import os
import re
import subprocess

from . import pannusch_conditional_tail_delivery as source_geometry
from . import pannusch_mass_delivery as source

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / 'docs/analysis/sci_md_trigonelline_delivery_001'
TASK = 'SCI-MD-TRIGONELLINE-DELIVERY-001'
SPECIES = 'trigonelline'
ARMS = ('TR-K0', 'TR-D0')
SUFFIX = (3, 5, 7, 10)
ASSAY_FRACTIONS = (1, 2, 3, 5, 7, 10)
PANELS = source_geometry.PANELS
read, write, digest = source_geometry.read, source_geometry.write, source.digest
canonical, identity = source_geometry.md.canonical, source_geometry.md.identity
private_directory = source_geometry.private_directory


def git(expression):
    return subprocess.check_output(['git', 'rev-parse', expression], cwd=ROOT, text=True).strip()


def contract():
    value = read(DOC / 'INFORMATION_CONTRACT.json')
    if value['task'] != TASK or value['arms'] != list(ARMS):
        raise ValueError('WRONG_TRIGONELLINE_CONTRACT')
    subprocess.check_call(['git', 'merge-base', '--is-ancestor', value['contract_commit'], 'HEAD'], cwd=ROOT)
    for relative, sha in value['files'].items():
        committed = subprocess.check_output(['git', 'show', value['contract_commit'] + ':' + relative], cwd=ROOT)
        if digest(ROOT / relative) != sha or source.hashlib.sha256(committed).hexdigest() != sha:
            raise ValueError('PROTOCOL_MUST_BE_COMMITTED_AND_IMMUTABLE_BEFORE_FITTING')
    return value


def qualified_sources():
    history = read(DOC / 'SOURCE_IDENTITIES.json')
    source.verify_registers(history)
    paths = source.source_paths()
    root = os.environ.get('PUCKWORKS_EXTERNAL_DATA_ROOT')
    if not root:
        root = read(source.config_path())['sources']['PANNUSCH2024_MENDELEY_FULL_REPOSITORY']['path']
    for row in source.rows(source.DATA / 'source_inputs.csv'):
        if row['source_id'] in ('P24-HPLC-FIT', 'P24-HPLC-PRED'):
            paths[row['source_id']] = Path(root) / row['source_relpath']
    if {k: digest(v) for k, v in paths.items()} != history['source_files']:
        raise ValueError('TRIGONELLINE_SOURCE_IDENTITY_MISMATCH')
    for relative, sha in read(DOC / 'DEPENDENCIES.json')['files'].items():
        if digest(ROOT / relative) != sha:
            raise ValueError('ACCEPTED_DEPENDENCY_CHANGED:' + relative)
    return paths, history


def mass_projection(coords):
    """Project measured early vial masses, without reading any chemical column."""
    result = {}
    for shot in sorted({c['shot'] for c in coords}):
        early = sorted((c for c in coords if c['shot'] == shot and c['fraction'] in (1, 2)),
                       key=lambda c: c['fraction'])
        if (len(early) != 2 or any(c['coordinate_status'] != 'QUALIFIED' for c in early)
                or early[0]['b0'] != 0 or early[0]['b1'] != early[1]['b0']):
            raise ValueError('QUALIFIED_ORIGINAL_EARLY_VIAL_MASSES_REQUIRED')
        masses = [source_geometry.md.number(c['mass_kg']) for c in early]
        if any(m <= 0 for m in masses):
            raise ValueError('POSITIVE_MEASURED_VIAL_MASSES_REQUIRED')
        result[shot] = masses
    return result


def analyte_index(rows, campaign):
    selected = [r for r in rows if r['campaign_id'] == campaign
                and r['analyte'] == SPECIES and int(r['fraction_id']) in SUFFIX]
    index = {(r['shot_id'], int(r['fraction_id'])): r for r in selected}
    if len(index) != len(selected):
        raise ValueError('DUPLICATE_TRIGONELLINE_SOURCE_IDENTITY')
    return index


def source_value(row, campaign):
    if row['concentration_unit'] != 'mg/g':
        raise ValueError('TRIGONELLINE_SOURCE_UNIT_MISMATCH')
    field = 'concentration_value' if campaign == 'FIT_2021_12' else 'measured_concentration'
    q = source_geometry.md.number(float(row[field])) / 1000
    if q < 0:
        raise ValueError('NEGATIVE_TRIGONELLINE_REQUIRES_SOURCE_ADJUDICATION')
    return q


def target_slots(coords, rows, *, campaign, include_values=False):
    if campaign not in ('FIT_2021_12', 'PREDICTION_2022_03'):
        raise ValueError('DECLARED_CAMPAIGN_REQUIRED')
    if include_values and campaign != 'FIT_2021_12':
        raise ValueError('PRED_VALUES_REQUIRE_THE_SINGLE_APPROVED_OUTCOME_JOIN')
    index = analyte_index(rows, campaign)
    slots = []
    for c in coords:
        if c['campaign'] != campaign or c['fraction'] not in SUFFIX:
            continue
        r = index.get((c['shot'], c['fraction']))
        field = 'concentration_value' if campaign == 'FIT_2021_12' else 'measured_concentration'
        if r is None:
            reason = 'MISSING_TRIGONELLINE_SOURCE_ROW'
        else:
            if r['concentration_unit'] != 'mg/g':
                raise ValueError('TRIGONELLINE_SOURCE_UNIT_MISMATCH')
            reason = (r.get('exclusion_reason') or 'INVALID_TRIGONELLINE'
                      if r['validity'] != 'VALID' or r.get('exclusion_reason') else
                      'MISSING_TRIGONELLINE_VALUE' if not r.get(field) else '')
            if not reason:
                source_value(r, campaign)  # Silent finite/nonnegative metadata qualification.
        _, exp, rep = c['shot'].split('-')
        item = {k: c[k] for k in ('campaign', 'condition', 'shot', 'fraction', 'b0', 'b1',
                                  'mass_kg', 'coordinate_status')}
        item.update(analyte=SPECIES, source_id=r['source_id'] if r else None,
                    source_field=f'ExperimentalData({int(exp[1:])}).run({int(rep[1:])}).cAlcaloids({ASSAY_FRACTIONS.index(c["fraction"])+1},2)',
                    analyte_eligible=not reason, source_reason=reason,
                    validity_reason=reason or 'VALID_SOURCE_TRIGONELLINE')
        if include_values:
            item['q'] = source_value(r, campaign) if not reason else None
        slots.append(item)
    return slots


def original_reconciliation(paths, slots, rows):
    """Silent source check; emits counts only, including for PRED."""
    objects, books, count = {}, {}, 0
    indices = {c: analyte_index(rows, c) for c in ('FIT_2021_12', 'PREDICTION_2022_03')}
    for slot in slots:
        if not slot['analyte_eligible']:
            continue
        label, exp, rep = slot['shot'].split('-')
        if label not in objects:
            objects[label] = source.mat(paths['P24-MAT-' + label])['ExperimentalData']
            books[label] = source.workbook(paths['P24-HPLC-' + label])
        e, j = int(exp[1:]), int(rep[1:])
        k = ASSAY_FRACTIONS.index(slot['fraction'])
        run = objects[label][e-1].run[j-1]
        original = source_geometry.md.number(float(run.cAlcaloids[k, 1]))
        row = indices[slot['campaign']][slot['shot'], slot['fraction']]
        q = source_value(row, slot['campaign'])
        mass_mg = books[label][f'{3*e-2}-{3*e}'].cell(10+6*(j-1)+k, 20).value
        if (original < 0 or abs(original/1000-q) > 5.001e-12
                or abs(float(mass_mg)/float(run.mE[slot['fraction']-1])-original) > 1e-12):
            raise ValueError('ORIGINAL_TRIGONELLINE_RECONSTRUCTION_MISMATCH')
        count += 1
    return count


def hplc_formula_check(paths):
    import openpyxl
    result = {}
    for label in ('FIT', 'PRED'):
        path = paths['P24-HPLC-' + label]
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        count = 0
        sign, offset, slope = ('+', .3151, 9.1718) if label == 'FIT' else ('-', 1.9691, 9.4635)
        for sheet in formulas.sheetnames:
            if not re.fullmatch(r'\d+-\d+', sheet):
                continue
            fs, cs = list(formulas[sheet].iter_rows()), list(cached[sheet].iter_rows())
            if fs[8][19].value != 'Trigonellin':
                raise ValueError('ORIGINAL_ANALYTE_HEADER_MISMATCH')
            for rep in range(3):
                for k in (2, 3, 4, 5):
                    row = 10+6*rep+k
                    formula = fs[row-1][19].value
                    match = re.match(r'=\(D(\d+)', str(formula))
                    if not match:
                        raise ValueError('UNRESOLVED_TRIGONELLINE_CALIBRATION')
                    origin = int(match[1])
                    if formula != f'=(D{origin}{sign}{offset})/({slope}*1000)*$C{origin}*$B{origin}':
                        raise ValueError('TRIGONELLINE_DILUTION_FORMULA_CHANGED')
                    mass, dilution, area = [source_geometry.md.number(cs[origin-1][col].value) for col in (1, 2, 3)]
                    expected = (area+(offset if label == 'FIT' else -offset))/(slope*1000)*dilution*mass
                    actual = source_geometry.md.number(cs[row-1][19].value)
                    if mass <= 0 or dilution <= 0 or expected < 0 or abs(expected-actual) > 1e-12:
                        raise ValueError('TRIGONELLINE_ORIGINAL_CALIBRATION_RECONCILIATION_FAILED')
                    count += 1
        formulas.close(); cached.close()
        result[label] = {'qualified_target_cells': count, 'status': 'PASS'}
    if result['FIT']['qualified_target_cells'] != 180 or result['PRED']['qualified_target_cells'] != 96:
        raise ValueError('ORIGINAL_HPLC_TARGET_MATRIX_REQUIRED')
    return result


def build_projections(coords, rows, groups):
    early = mass_projection(coords)
    fit_slots = target_slots(coords, rows, campaign='FIT_2021_12', include_values=True)
    queries = target_slots(coords, rows, campaign='PREDICTION_2022_03')
    training = []
    for shot in sorted(s for s in early if s.startswith('FIT-')):
        slots = [s for s in fit_slots if s['shot'] == shot]
        windows = []
        for s in slots:
            reason = s['source_reason'] or (s['coordinate_status'] if s['coordinate_status'] != 'QUALIFIED' else '')
            s.update(supported=not reason, reason=reason)
            if not reason:
                anchor = type('Anchor', (), {'b_anchor': sum(early[shot])})()
                start, allowance = source_geometry.source_query_start(s, anchor)
                s.update(integration_start_kg=start, coordinate_allowance_kg=allowance)
                windows.append({'fraction': s['fraction'], 'start_kg': start, 'end_kg': s['b1'],
                                'q': s['q'], 'mass_kg': s['mass_kg'], 'coordinate_allowance_kg': allowance})
        if not windows:
            raise ValueError('ORIGINAL_FIT_SHOT_WITHOUT_SUPERVISION')
        training.append({'shot': shot, 'group': groups[slots[0]['condition']],
                         'early_values': early[shot], 'windows': windows})
    domain = max(w['end_kg'] for s in training for w in s['windows'])
    for q in queries:
        reason = q['source_reason'] or (q['coordinate_status'] if q['coordinate_status'] != 'QUALIFIED' else
                 'OUTSIDE_COMMON_TRAINING_MASS_DOMAIN' if q['b1'] > domain else '')
        q.update(supported=not reason, reason=reason)
    originals = sorted({s['group'] for s in training})
    if len(training) != 45 or len(originals) != 15 or len(fit_slots) != 180 or len(queries) != 96:
        raise ValueError('ORIGINAL_SOURCE_DENOMINATORS_REQUIRED')
    for group in originals:
        if sum(len(s['windows']) >= 3 for s in training if s['group'] == group) < 2:
            raise ValueError('TWO_PHYSICAL_SHOTS_WITH_THREE_TARGET_WINDOWS_REQUIRED')
    primary = [q for q in queries if q['condition'] in PANELS['primary']]
    if len(primary) != 48 or len({q['shot'] for q in primary}) != 12 or not all(q['supported'] for q in primary):
        raise ValueError('ALL_FORTY_EIGHT_PRIMARY_WINDOWS_REQUIRED')
    folds = []
    for group in originals:
        retained = [s for s in training if s['group'] != group]
        held = [s for s in training if s['group'] == group]
        bound = max(w['end_kg'] for s in retained for w in s['windows'])
        support = [(s['shot'], w['fraction']) for s in held for w in s['windows'] if w['end_kg'] <= bound]
        held_ids = {s['shot'] for s in held}
        folds.append({'group': group, 'domain_kg': bound, 'training_shots': [s['shot'] for s in retained],
                      'held_shots': [s['shot'] for s in held], 'support': support,
                      'slots': [{'shot': s['shot'], 'fraction': s['fraction'],
                                 'supported': (s['shot'], s['fraction']) in support,
                                 'reason': s['reason'] or ('OUTSIDE_FOLD_DOMAIN' if (s['shot'], s['fraction']) not in support else '')}
                                for s in fit_slots if s['shot'] in held_ids]})
    counts = {'fit_physical_shots': 45, 'fit_original_designs': 15, 'fit_intended_slots': 180,
              'fit_supported_slots': sum(s['supported'] for s in fit_slots),
              'fit_reasons': dict(Counter(s['reason'] for s in fit_slots if s['reason'])),
              'common_domain_kg': [0., domain], 'PRED_values_projected': 0,
              'panels': {p: {'original_shots': 3*len(cs), 'intended_slots': 12*len(cs),
                             'supported_slots': sum(q['supported'] for q in queries if q['condition'] in cs)}
                         for p, cs in PANELS.items()}}
    clean_coords = [{k: c[k] for k in ('campaign', 'condition', 'shot', 'fraction', 'b0', 'b1',
                                     'mass_kg', 'coordinate_status')} for c in coords]
    return {'training': training, 'fit_slots': fit_slots, 'queries': queries,
            'early_inputs': {k: v for k, v in early.items() if k.startswith('PRED-')},
            'development_support': folds, 'source_counts': counts, 'coordinates': clean_coords}


def prepare(out):
    authority = contract()
    paths, history = qualified_sources()
    out = private_directory(out, create=True)
    out.chmod(0o700)
    coords, rows = source_geometry.coordinates(paths), source_geometry.all_rows()
    projections = build_projections(coords, rows, source_geometry.condition_groups())
    checked = original_reconciliation(paths, projections['fit_slots']+projections['queries'], rows)
    if checked != 276:
        raise ValueError('ALL_ORIGINAL_TRIGONELLINE_TARGETS_MUST_RECONCILE')
    projections['source_counts']['original_reconciliations'] = checked
    projections['source_counts']['HPLC_formula_check'] = hplc_formula_check(paths)
    projections.update(source=dict(history, rights=source.RIGHTS), information_contract=authority)
    for name, value in projections.items():
        write(out / (name+'.json'), value)
    write(out / 'prepared_hashes.json', {str(p.relative_to(out)): digest(p) for p in out.glob('*.json')})
    write(out / 'preparation_runtime.json', source_geometry.imported_files())
    print(canonical(projections['source_counts']))


def main():
    parser = argparse.ArgumentParser(description=__doc__, allow_abbrev=False)
    parser.add_argument('command', choices=('prepare',))
    parser.add_argument('--out', required=True, type=Path)
    args = parser.parse_args()
    prepare(args.out)


if __name__ == '__main__':
    main()
