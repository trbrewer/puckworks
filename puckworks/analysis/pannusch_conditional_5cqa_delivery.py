"""Task-local five_cqa projection and frozen research comparison.

Source originals and row artifacts stay private. Preparation exposes only FIT
chemistry; PRED chemistry crosses the boundary only in the approved scorer.
"""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
import os
import platform
import re
import subprocess
import time

import numpy as np
import scipy

from . import pannusch_conditional_tail_delivery as source_geometry
from . import pannusch_mass_delivery as source
from . import conditional_5cqa_delivery as md

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / 'docs/analysis/sci_md_5cqa_delivery_001'
TASK = 'SCI-MD-5CQA-DELIVERY-001'
SPECIES = '5CQA'
ARMS = ('E0', 'D0')
SUFFIX = (3, 5, 7, 10)
ASSAY_FRACTIONS = (1, 2, 3, 5, 7, 10)
PANELS = source_geometry.PANELS
read, write, digest = source_geometry.read, source_geometry.write, source.digest
canonical, identity = source_geometry.md.canonical, source_geometry.md.identity
private_directory = source_geometry.private_directory
TASK_DEADLINE = datetime.fromisoformat('2026-09-29T04:48:20+00:00').timestamp()


def check_task_clock():
    if time.time() >= TASK_DEADLINE:
        raise RuntimeError('SIX_HOUR_TASK_DEADLINE')


def git(expression):
    return subprocess.check_output(['git', 'rev-parse', expression], cwd=ROOT, text=True).strip()


def contract():
    value = read(DOC / 'INFORMATION_CONTRACT.json')
    if value['task'] != TASK or value['arms'] != list(ARMS):
        raise ValueError('WRONG_FIVE_CQA_CONTRACT')
    subprocess.check_call(['git', 'merge-base', '--is-ancestor', value['contract_commit'], 'HEAD'], cwd=ROOT)
    for relative, sha in value['files'].items():
        committed = subprocess.check_output(['git', 'show', value['contract_commit'] + ':' + relative], cwd=ROOT)
        if digest(ROOT / relative) != sha or source.hashlib.sha256(committed).hexdigest() != sha:
            raise ValueError('PROTOCOL_MUST_BE_COMMITTED_AND_IMMUTABLE_BEFORE_FITTING')
    return value


def qualified_sources():
    history = read(DOC / 'SOURCE_IDENTITIES.json')
    source.verify_registers(history)
    paths = source.source_paths()
    root = os.environ.get('PUCKWORKS_EXTERNAL_DATA_ROOT')
    if not root:
        root = read(source.config_path())['sources']['PANNUSCH2024_MENDELEY_FULL_REPOSITORY']['path']
    for row in source.rows(source.DATA / 'source_inputs.csv'):
        if row['source_id'] in ('P24-HPLC-FIT', 'P24-HPLC-PRED'):
            paths[row['source_id']] = Path(root) / row['source_relpath']
    if {k: digest(v) for k, v in paths.items()} != history['source_files']:
        raise ValueError('FIVE_CQA_SOURCE_IDENTITY_MISMATCH')
    for relative, sha in read(DOC / 'DEPENDENCIES.json')['files'].items():
        if digest(ROOT / relative) != sha:
            raise ValueError('ACCEPTED_DEPENDENCY_CHANGED:' + relative)
    return paths, history


def mass_projection(coords):
    """Project measured early vial masses, without reading any chemical column."""
    result = {}
    for shot in sorted({c['shot'] for c in coords}):
        early = sorted((c for c in coords if c['shot'] == shot and c['fraction'] in (1, 2)),
                       key=lambda c: c['fraction'])
        if (len(early) != 2 or any(c['coordinate_status'] != 'QUALIFIED' for c in early)
                or early[0]['b0'] != 0 or early[0]['b1'] != early[1]['b0']):
            raise ValueError('QUALIFIED_ORIGINAL_EARLY_VIAL_MASSES_REQUIRED')
        masses = [source_geometry.md.number(c['mass_kg']) for c in early]
        if any(m <= 0 for m in masses):
            raise ValueError('POSITIVE_MEASURED_VIAL_MASSES_REQUIRED')
        result[shot] = masses
    return result


def analyte_index(rows, campaign):
    selected = [r for r in rows if r['campaign_id'] == campaign
                and r['analyte'] == SPECIES and int(r['fraction_id']) in SUFFIX]
    index = {(r['shot_id'], int(r['fraction_id'])): r for r in selected}
    if len(index) != len(selected):
        raise ValueError('DUPLICATE_FIVE_CQA_SOURCE_IDENTITY')
    return index


def source_value(row, campaign):
    if row['concentration_unit'] != 'mg/g':
        raise ValueError('FIVE_CQA_SOURCE_UNIT_MISMATCH')
    field = 'concentration_value' if campaign == 'FIT_2021_12' else 'measured_concentration'
    q = source_geometry.md.number(float(row[field])) / 1000
    if q < 0:
        raise ValueError('NEGATIVE_FIVE_CQA_REQUIRES_SOURCE_ADJUDICATION')
    return q


def target_slots(coords, rows, *, campaign, include_values=False):
    if campaign not in ('FIT_2021_12', 'PREDICTION_2022_03'):
        raise ValueError('DECLARED_CAMPAIGN_REQUIRED')
    if include_values and campaign != 'FIT_2021_12':
        raise ValueError('PRED_VALUES_REQUIRE_THE_SINGLE_APPROVED_OUTCOME_JOIN')
    index = analyte_index(rows, campaign)
    slots = []
    for c in coords:
        if c['campaign'] != campaign or c['fraction'] not in SUFFIX:
            continue
        r = index.get((c['shot'], c['fraction']))
        field = 'concentration_value' if campaign == 'FIT_2021_12' else 'measured_concentration'
        if r is None:
            reason = 'MISSING_FIVE_CQA_SOURCE_ROW'
        else:
            if r['concentration_unit'] != 'mg/g':
                raise ValueError('FIVE_CQA_SOURCE_UNIT_MISMATCH')
            reason = (r.get('exclusion_reason') or 'INVALID_FIVE_CQA'
                      if r['validity'] != 'VALID' or r.get('exclusion_reason') else
                      'MISSING_FIVE_CQA_VALUE' if not r.get(field) else '')
            if not reason:
                source_value(r, campaign)  # Silent finite/nonnegative metadata qualification.
        _, exp, rep = c['shot'].split('-')
        item = {k: c[k] for k in ('campaign', 'condition', 'shot', 'fraction', 'b0', 'b1',
                                  'mass_kg', 'coordinate_status')}
        item.update(analyte=SPECIES, source_id=r['source_id'] if r else None,
                    source_field=f'ExperimentalData({int(exp[1:])}).run({int(rep[1:])}).cAlcaloids({ASSAY_FRACTIONS.index(c["fraction"])+1},3)',
                    analyte_eligible=not reason, source_reason=reason,
                    validity_reason=reason or 'VALID_SOURCE_FIVE_CQA')
        if include_values:
            item['q'] = source_value(r, campaign) if not reason else None
        slots.append(item)
    return slots


def original_reconciliation(paths, slots, rows):
    """Silent source check; emits counts only, including for PRED."""
    objects, books, count = {}, {}, 0
    indices = {c: analyte_index(rows, c) for c in ('FIT_2021_12', 'PREDICTION_2022_03')}
    for slot in slots:
        if not slot['analyte_eligible']:
            continue
        label, exp, rep = slot['shot'].split('-')
        if label not in objects:
            objects[label] = source.mat(paths['P24-MAT-' + label])['ExperimentalData']
            books[label] = source.workbook(paths['P24-HPLC-' + label])
        e, j = int(exp[1:]), int(rep[1:])
        k = ASSAY_FRACTIONS.index(slot['fraction'])
        run = objects[label][e-1].run[j-1]
        original = source_geometry.md.number(float(run.cAlcaloids[k, 2]))
        row = indices[slot['campaign']][slot['shot'], slot['fraction']]
        q = source_value(row, slot['campaign'])
        mass_mg = source_geometry.md.number(books[label][f'{3*e-2}-{3*e}'].cell(10+6*(j-1)+k, 25).value)
        if (original < 0 or abs(original/1000-q) > 5.001e-12
                or abs(float(row['fraction_liquid_g_or_ml'])-float(run.mE[slot['fraction']-1])) > 5.001e-9
                or abs(float(mass_mg)/float(run.mE[slot['fraction']-1])-original) > 1e-12):
            raise ValueError('ORIGINAL_FIVE_CQA_RECONSTRUCTION_MISMATCH')
        count += 1
    return count


def hplc_formula_check(paths):
    import openpyxl
    result = {}
    for label in ('FIT', 'PRED'):
        path = paths['P24-HPLC-' + label]
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        count = 0
        sign, offset, slope = ('-', 78.923, 24.513) if label == 'FIT' else ('+', 88.067, 26.383)
        for sheet in formulas.sheetnames:
            if not re.fullmatch(r'\d+-\d+', sheet):
                continue
            fs, cs = list(formulas[sheet].iter_rows()), list(cached[sheet].iter_rows())
            if fs[8][24].value != 'Chlorogenic acid':
                raise ValueError('ORIGINAL_ANALYTE_HEADER_MISMATCH')
            for rep in range(3):
                for k in (2, 3, 4, 5):
                    row = 10+6*rep+k
                    formula = fs[row-1][24].value
                    match = re.match(r'=\(I(\d+)', str(formula))
                    if not match:
                        raise ValueError('UNRESOLVED_FIVE_CQA_CALIBRATION')
                    origin = int(match[1])
                    if formula != f'=(I{origin}{sign}{offset})/({slope}*1000)*$C{origin}*$B{origin}':
                        raise ValueError('FIVE_CQA_DILUTION_FORMULA_CHANGED')
                    mass, dilution, area = [source_geometry.md.number(cs[origin-1][col].value) for col in (1, 2, 8)]
                    expected = (area+(-offset if label == 'FIT' else offset))/(slope*1000)*dilution*mass
                    actual = source_geometry.md.number(cs[row-1][24].value)
                    if mass <= 0 or dilution <= 0 or expected < 0 or abs(expected-actual) > 1e-12:
                        raise ValueError('FIVE_CQA_ORIGINAL_CALIBRATION_RECONCILIATION_FAILED')
                    count += 1
        formulas.close(); cached.close()
        result[label] = {'qualified_target_cells': count, 'status': 'PASS'}
    if result['FIT']['qualified_target_cells'] != 180 or result['PRED']['qualified_target_cells'] != 96:
        raise ValueError('ORIGINAL_HPLC_TARGET_MATRIX_REQUIRED')
    return result


def original_rows(paths, rows, campaign):
    """Attach unrounded original chemistry only for the explicitly permitted role.

    The public table's rounding is reconciled separately and is not uncertainty.
    PRED callers are confined to the exclusive, approved score stage.
    """
    if campaign not in ('FIT_2021_12', 'PREDICTION_2022_03'):
        raise ValueError('DECLARED_CHEMISTRY_ROLE_REQUIRED')
    label = 'FIT' if campaign == 'FIT_2021_12' else 'PRED'
    field = 'concentration_value' if label == 'FIT' else 'measured_concentration'
    objects = source.mat(paths['P24-MAT-'+label])['ExperimentalData']
    result = []
    for row in rows:
        r = dict(row)
        if (r['campaign_id'] == campaign and r['analyte'] == SPECIES
                and int(r['fraction_id']) in SUFFIX and r['validity'] == 'VALID'
                and not r.get('exclusion_reason') and r.get(field)):
            _, exp, rep = r['shot_id'].split('-')
            k = ASSAY_FRACTIONS.index(int(r['fraction_id']))
            value = md.number(float(objects[int(exp[1:])-1].run[int(rep[1:])-1].cAlcaloids[k, 2]))
            if value < 0 or abs(value-float(r[field])) > 5.001e-9:
                raise ValueError('ORIGINAL_FIVE_CQA_EXPORT_ROUNDING_MISMATCH')
            r[field] = repr(value)
        result.append(r)
    return result


def build_projections(coords, rows, groups):
    early = mass_projection(coords)
    fit_slots = target_slots(coords, rows, campaign='FIT_2021_12', include_values=True)
    queries = target_slots(coords, rows, campaign='PREDICTION_2022_03')
    training = []
    for shot in sorted(s for s in early if s.startswith('FIT-')):
        slots = [s for s in fit_slots if s['shot'] == shot]
        windows = []
        for s in slots:
            reason = s['source_reason'] or (s['coordinate_status'] if s['coordinate_status'] != 'QUALIFIED' else '')
            s.update(supported=not reason, reason=reason)
            if not reason:
                anchor = type('Anchor', (), {'b_anchor': sum(early[shot])})()
                start, allowance = source_geometry.source_query_start(s, anchor)
                s.update(integration_start_kg=start, coordinate_allowance_kg=allowance)
                windows.append({'fraction': s['fraction'], 'start_kg': start, 'end_kg': s['b1'],
                                'q': s['q'], 'mass_kg': s['mass_kg'], 'coordinate_allowance_kg': allowance})
        if not windows:
            raise ValueError('ORIGINAL_FIT_SHOT_WITHOUT_SUPERVISION')
        training.append({'shot': shot, 'group': groups[slots[0]['condition']],
                         'early_values': early[shot], 'windows': windows})
    domain = max(w['end_kg'] for s in training for w in s['windows'])
    for q in queries:
        reason = q['source_reason'] or (q['coordinate_status'] if q['coordinate_status'] != 'QUALIFIED' else
                 'OUTSIDE_COMMON_TRAINING_MASS_DOMAIN' if q['b1'] > domain else '')
        q.update(supported=not reason, reason=reason)
    originals = sorted({s['group'] for s in training})
    if len(training) != 45 or len(originals) != 15 or len(fit_slots) != 180 or len(queries) != 96:
        raise ValueError('ORIGINAL_SOURCE_DENOMINATORS_REQUIRED')
    for group in originals:
        if sum(len(s['windows']) >= 3 for s in training if s['group'] == group) < 2:
            raise ValueError('TWO_PHYSICAL_SHOTS_WITH_THREE_TARGET_WINDOWS_REQUIRED')
    primary = [q for q in queries if q['condition'] in PANELS['primary']]
    if len(primary) != 48 or len({q['shot'] for q in primary}) != 12 or not all(q['supported'] for q in primary):
        raise ValueError('ALL_FORTY_EIGHT_PRIMARY_WINDOWS_REQUIRED')
    folds = []
    for group in originals:
        retained = [s for s in training if s['group'] != group]
        held = [s for s in training if s['group'] == group]
        bound = max(w['end_kg'] for s in retained for w in s['windows'])
        support = [(s['shot'], w['fraction']) for s in held for w in s['windows'] if w['end_kg'] <= bound]
        held_ids = {s['shot'] for s in held}
        folds.append({'group': group, 'domain_kg': bound, 'training_shots': [s['shot'] for s in retained],
                      'held_shots': [s['shot'] for s in held], 'support': support,
                      'slots': [{'shot': s['shot'], 'fraction': s['fraction'],
                                 'supported': (s['shot'], s['fraction']) in support,
                                 'reason': s['reason'] or ('OUTSIDE_FOLD_DOMAIN' if (s['shot'], s['fraction']) not in support else '')}
                                for s in fit_slots if s['shot'] in held_ids]})
    counts = {'fit_physical_shots': 45, 'fit_original_designs': 15, 'fit_intended_slots': 180,
              'fit_supported_slots': sum(s['supported'] for s in fit_slots),
              'fit_reasons': dict(Counter(s['reason'] for s in fit_slots if s['reason'])),
              'common_domain_kg': [0., domain], 'PRED_values_projected': 0,
              'panels': {p: {'original_shots': 3*len(cs), 'intended_slots': 12*len(cs),
                             'supported_slots': sum(q['supported'] for q in queries if q['condition'] in cs)}
                         for p, cs in PANELS.items()}}
    clean_coords = [{k: c[k] for k in ('campaign', 'condition', 'shot', 'fraction', 'b0', 'b1',
                                     'mass_kg', 'coordinate_status')} for c in coords]
    return {'training': training, 'fit_slots': fit_slots, 'queries': queries,
            'early_inputs': {k: v for k, v in early.items() if k.startswith('PRED-')},
            'development_support': folds, 'source_counts': counts, 'coordinates': clean_coords}


def prepare(out):
    check_task_clock()
    authority = contract()
    paths, history = qualified_sources()
    out = private_directory(out, create=True)
    out.chmod(0o700)
    coords, rows = source_geometry.coordinates(paths), source_geometry.all_rows()
    projections = build_projections(coords, original_rows(paths, rows, 'FIT_2021_12'),
                                    source_geometry.condition_groups())
    counts = projections['source_counts']
    preflight = read(DOC / 'DATA_AVAILABILITY_PREFLIGHT.json')
    if (counts['fit_supported_slots'] != preflight['counts']['FIT']['supported_slots']
            or counts['common_domain_kg'] != preflight['training_domain_kg']
            or any(counts['panels'][p]['supported_slots'] != c['supported_slots']
                   for p, c in preflight['panels'].items())):
        raise ValueError('QUALIFIED_SOURCE_COHORT_OR_DOMAIN_CHANGED')
    checked = original_reconciliation(paths, projections['fit_slots']+projections['queries'], rows)
    if checked != 276:
        raise ValueError('ALL_ORIGINAL_FIVE_CQA_TARGETS_MUST_RECONCILE')
    projections['source_counts']['original_reconciliations'] = checked
    projections['source_counts']['HPLC_formula_check'] = hplc_formula_check(paths)
    projections.update(source=dict(history, rights=source.RIGHTS), information_contract=authority)
    for name, value in projections.items():
        write(out / (name+'.json'), value)
    write(out / 'prepared_hashes.json', {str(p.relative_to(out)): digest(p) for p in out.glob('*.json')})
    write(out / 'preparation_runtime.json', source_geometry.imported_files())
    print(canonical(projections['source_counts']))


def prepared(out):
    out = private_directory(out)
    contract(); qualified_sources()
    if any((out / n).exists() for n in ('freeze.json', 'score_receipt.json')):
        raise ValueError('FROZEN_RUN_IMMUTABLE')
    for relative, sha in read(out / 'prepared_hashes.json').items():
        if digest(out / relative) != sha:
            raise ValueError('PREPARED_PROJECTION_DRIFT')
    for relative, sha in read(out / 'preparation_runtime.json').items():
        if digest(ROOT / relative) != sha:
            raise ValueError('PREPARATION_RUNTIME_DRIFT')
    return out


def source_identity(out):
    value = read(out / 'source.json')
    return identity({'source_files': {k: v for k, v in value['source_files'].items() if k.endswith('-FIT')},
                     'fit_register': value['registers']['fit_fraction_replicates.csv']})


def fitting_budget(out):
    from . import conditional_5cqa_training as train
    if any(os.environ.get(k) != '1' for k in ('OPENBLAS_NUM_THREADS', 'OMP_NUM_THREADS', 'MKL_NUM_THREADS')):
        raise ValueError('ONE_BLAS_THREAD_REQUIRED')
    return train.Budget(out / 'starts', evidence_root=out.parent)


def develop_folds(out, budget):
    from . import conditional_5cqa_training as train
    out = prepared(out)
    dest = out / 'development'; dest.mkdir()
    records, specs = read(out / 'training.json'), read(out / 'development_support.json')
    result = {'task': TASK, 'scope': 'DEVELOPMENT_NOT_UNBIASED_NESTED_VALIDATION', 'arms': {}}
    for arm in ARMS:
        shots = train.project_arm(records, arm)
        candidates = []
        for lam in (None,) if arm == 'E0' else train.LAMBDAS:
            candidate = {'lambda': lam, 'status': 'SELECTABLE', 'folds': []}
            for spec in specs:
                group = spec['group']
                training, held = train.split_design(shots, group)
                support = {tuple(v) for v in spec['support']}
                expected = {(s.shot, w.fraction) for s in held for w in s.windows if w.end_kg <= spec['domain_kg']}
                if ([s.shot for s in training] != spec['training_shots']
                        or [s.shot for s in held] != spec['held_shots'] or support != expected
                        or max(w.end_kg for s in training for w in s.windows) != spec['domain_kg']):
                    raise ValueError('WHOLE_ORIGINAL_DESIGN_FOLD_OR_SUPPORT_DRIFT')
                name = f'{arm}.{lam}.{group}'
                model, audit = train.fit(training, lam, budget, name, source_identity(out), source.RIGHTS)
                write(dest / (name+'.audit.json'), audit)
                fold = {'group': group, 'status': audit.get('status', 'CONVERGED'),
                        'domain_kg': spec['domain_kg'], 'training_shots': spec['training_shots'],
                        'held_shots': spec['held_shots'], 'support': spec['support']}
                if model is None:
                    candidate['status'] = 'NONSELECTABLE'
                else:
                    model.save(dest / (name+'.model.json'))
                    value, allowance = train.held_metrics(model, held, support)
                    fold.update(R_mg_g=value, max_allowance_kg=allowance, model_sha256=model.sha256)
                candidate['folds'].append(fold)
                write(dest / (name+'.fold.json'), fold)
            candidate['balanced_R_mg_g'] = (float(np.mean([f['R_mg_g'] for f in candidate['folds']]))
                                            if candidate['status'] == 'SELECTABLE' else None)
            candidates.append(candidate)
        result['arms'][arm] = {'candidates': candidates,
                               'selected_lambda': train.select_lambda(candidates) if arm == 'D0' else None}
    write(dest / 'development.json', result)
    print(canonical({arm: {'selected_lambda': r['selected_lambda'],
                           'candidates': [{k: c[k] for k in ('lambda', 'status', 'balanced_R_mg_g')} for c in r['candidates']]}
                     for arm, r in result['arms'].items()}))


def final_fit(out, budget):
    from . import conditional_5cqa_training as train
    out = prepared(out)
    dest = out / 'models'; dest.mkdir()
    selection = read(out / 'development/development.json')
    records = read(out / 'training.json')
    for arm in ARMS:
        shots = train.project_arm(records, arm)
        model, audit = train.fit(shots, selection['arms'][arm]['selected_lambda'], budget,
                                 arm+'.final', source_identity(out), source.RIGHTS, scope='FINAL_FIT')
        write(dest / (arm+'.audit.json'), audit)
        if model is None:
            raise ValueError('FINAL_FIVE_CQA_MODEL_NONSELECTABLE')
        model.save(dest / (arm+'.json'))
        all_windows = {(s.shot, w.fraction) for s in shots for w in s.windows}
        value, allowance = train.held_metrics(model, shots, all_windows)
        write(dest / (arm+'.qualification.json'), {'status': 'PASS', 'FIT_only_R_mg_g': value,
                                                  'max_allowance_kg': allowance, 'model_sha256': model.sha256})
    print('FINAL_MODELS_NUMERICALLY_QUALIFIED')


def develop(out):
    check_task_clock()
    out = prepared(out)
    budget = fitting_budget(out)
    try:
        develop_folds(out, budget)
        final_fit(out, budget)
        write(out / 'fitting_completion.json', {'status': 'COMPLETE',
              'fitting_stage_wall_seconds': time.time()-budget.clock['started_unix']})
    finally:
        budget.worker_lock.close()


QUERY_KEYS = ('campaign', 'condition', 'shot', 'fraction', 'b0', 'b1', 'mass_kg', 'coordinate_status',
              'analyte', 'source_id', 'source_field', 'analyte_eligible', 'source_reason',
              'validity_reason', 'supported', 'reason')


def predict_records(models, early, queries):
    if set(models) != set(ARMS):
        raise ValueError('BOTH_DECLARED_FIVE_CQA_MODELS_REQUIRED')
    predictions, states = {}, {}
    for query in queries:
        md.exact_keys(query, QUERY_KEYS)
        if query['analyte'] != SPECIES or query['campaign'] != 'PREDICTION_2022_03':
            raise ValueError('COORDINATE_ONLY_PRED_FIVE_CQA_QUERY_REQUIRED')
    for arm in ARMS:
        results = []
        for q in queries:
            result = dict(q, prediction_status=q['reason'] or 'NOT_PREDICTED', numerical_qualified=False,
                          five_cqa_kg=None, five_cqa_mg=None, five_cqa_mg_g=None,
                          allowance_kg=None, feature_extrapolation=[])
            if q['supported']:
                try:
                    state = models[arm].condition(md.EarlyInput(arm, early[q['shot']], 'SOURCE_EARLY_INPUT'))
                    start, coord_allowance = source_geometry.source_query_start(q, state)
                    p = state.predict_intervals([start], [q['b1']])[0]
                    allowance = p.allowance_kg+coord_allowance
                    result.update(asdict(p), integration_start_kg=start, coordinate_allowance_kg=coord_allowance,
                                  allowance_kg=allowance, numerical_qualified=p.numerical_qualified and allowance <= 1e-9,
                                  prediction_status='QUALIFIED' if p.numerical_qualified and allowance <= 1e-9 else 'NUMERICALLY_UNQUALIFIED')
                    states[arm+':'+q['shot']] = state.to_dict()
                except (ValueError, FloatingPointError) as exc:
                    result['prediction_status'] = 'PREDICTION_FAILED:'+str(exc)
            results.append(result)
        predictions[arm] = results
    return predictions, states


def execution_account(out):
    starts = list((out / 'starts').glob('*.start.json'))
    logs = [read(p) for p in (out / 'starts').glob('*.end.json')]
    return {'task': TASK, 'iterative_starts': len(starts), 'completed_starts': len(logs),
            'starts_by_arm': {arm: len(list((out / 'starts').glob(arm+'.*.start.json'))) for arm in ARMS},
            'fitting_stage_wall_seconds': read(out / 'fitting_completion.json')['fitting_stage_wall_seconds'],
            'actual_residual_calls': sum(r['actual_residual_calls'] for r in logs),
            'numerical_jacobian_residual_calls': sum(r['numerical_jacobian_residual_calls'] for r in logs),
            'max_calls_per_start': max((r['actual_residual_calls'] for r in logs), default=0),
            'optimizer_wall_seconds': sum(r['elapsed_seconds'] for r in logs),
            'failed_starts': sum(r['status'] != 'CONVERGED' for r in logs),
            'boundary_starts': sum(bool(r['boundary_indices']) for r in logs),
            'peak_rss_kib': max((r['peak_rss_kib'] for r in logs), default=0),
            'workers': 1, 'blas_threads': 1, 'native_ewp_runs': 0, 'scientific_score_passes': 0,
            'python': platform.python_version(), 'numpy': np.__version__, 'scipy': scipy.__version__,
            'imported_first_party_modules': source_geometry.imported_files()}


def predict(out):
    out = prepared(out)
    models = {arm: md.Model.load(out / 'models' / (arm+'.json')) for arm in ARMS}
    started = time.monotonic()
    predictions, states = predict_records(models, read(out / 'early_inputs.json'), read(out / 'queries.json'))
    write(out / 'predictions.json', predictions); write(out / 'states.json', states)
    execution = execution_account(out)
    execution.update(prediction_wall_seconds=time.monotonic()-started,
                     max_allowance_kg=max(p['allowance_kg'] for ps in predictions.values() for p in ps if p['allowance_kg'] is not None))
    write(out / 'execution.json', execution)
    for ps in predictions.values():
        if len(ps) != 96 or any(not p['numerical_qualified'] for p in ps if p['supported']):
            raise ValueError('COMPLETE_QUALIFIED_SUPPORTED_PRED_MATRIX_REQUIRED')
    print(canonical({k: v for k, v in execution.items() if k != 'imported_first_party_modules'}))


def freeze(out):
    check_task_clock()
    out = prepared(out)
    if subprocess.check_output(['git', 'status', '--porcelain'], cwd=ROOT, text=True).strip():
        raise ValueError('COMMITTED_CLEAN_PRODUCER_REQUIRED')
    if not (out / 'predictions.json').exists():
        predict(out)
    execution = read(out / 'execution.json')
    if (execution['iterative_starts'] > 231 or execution['iterative_starts'] != execution['completed_starts']
            or execution['starts_by_arm'] != {'E0': 48, 'D0': 183}
            or execution['max_calls_per_start'] > 8000 or execution['fitting_stage_wall_seconds'] >= 3600):
        raise ValueError('EXECUTION_ACCOUNTING_NOT_QUALIFIED')
    predictions = read(out / 'predictions.json')
    for ps in predictions.values():
        primary = [p for p in ps if p['condition'] in PANELS['primary']]
        if len(primary) != 48 or not all(p['supported'] and p['numerical_qualified'] for p in primary):
            raise ValueError('PRIMARY_FREEZE_MUST_HAVE_ALL_QUALIFIED_WINDOWS')
    files = set(execution['imported_first_party_modules']) | set(read(out / 'preparation_runtime.json'))
    files |= set(read(DOC / 'DEPENDENCIES.json')['files'])
    files |= {str(p.relative_to(ROOT)) for p in DOC.iterdir() if p.is_file() and p.name not in ('README.md',)}
    files |= {str(p.relative_to(ROOT)) for p in (DOC / 'models').glob('*.json')}
    files |= {str(p.relative_to(ROOT)) for p in (ROOT / 'tests').glob('test_*5cqa*.py')}
    files |= {'puckworks/analysis/conditional_5cqa_training.py'}
    write(out / 'freeze.json', {'task': TASK, 'producer_commit': git('HEAD'), 'producer_tree': git('HEAD^{tree}'),
        'code_and_protocol': {p: digest(ROOT / p) for p in sorted(files)},
        'artifacts': {str(p.relative_to(out)): digest(p) for p in sorted(out.rglob('*.json'))},
        'arms': list(ARMS), 'panels': PANELS, 'outcomes_attached': False,
        'review_status': 'READY_FOR_INDEPENDENT_REVIEW', 'scoring_policy': 'ONE_INDEPENDENTLY_APPROVED_SCORE_NO_RETUNING'})
    print('READY_FOR_INDEPENDENT_REVIEW', digest(out / 'freeze.json'))


def verify_frozen(out):
    frozen = read(out / 'freeze.json')
    if frozen['task'] != TASK or frozen['arms'] != list(ARMS) or frozen['outcomes_attached'] is not False:
        raise ValueError('WRONG_FROZEN_EXPERIMENT')
    for relative, sha in frozen['code_and_protocol'].items():
        if digest(ROOT / relative) != sha:
            raise ValueError('FROZEN_CODE_OR_PROTOCOL_DRIFT:' + relative)
        committed = subprocess.check_output(['git', 'show', frozen['producer_commit']+':'+relative], cwd=ROOT)
        if source.hashlib.sha256(committed).hexdigest() != sha:
            raise ValueError('FROZEN_PRODUCER_COMMIT_BINDING_MISMATCH')
    if git(frozen['producer_commit']+'^{tree}') != frozen['producer_tree']:
        raise ValueError('FROZEN_PRODUCER_TREE_MISMATCH')
    for relative, sha in frozen['artifacts'].items():
        if digest(out / relative) != sha:
            raise ValueError('FROZEN_ARTIFACT_DRIFT:' + relative)
    source.verify_registers(read(out / 'source.json'))
    return frozen


def verify(out):
    out = private_directory(out)
    frozen = verify_frozen(out)
    models = {arm: md.Model.load(out / 'models' / (arm+'.json')) for arm in ARMS}
    predictions, states = predict_records(models, read(out / 'early_inputs.json'), read(out / 'queries.json'))
    if canonical(predictions) != canonical(read(out / 'predictions.json')) or canonical(states) != canonical(read(out / 'states.json')):
        raise ValueError('SAVED_MODEL_INFERENCE_NOT_DETERMINISTIC')
    result = {'task': TASK, 'status': 'PASS', 'reviewed_producer': frozen['producer_commit'],
              'prediction_slots': sum(map(len, predictions.values())), 'prediction_max_difference': 0.,
              'optimizer_calls': 0, 'outcome_joins': 0, 'score_calls': 0, 'writes': 0}
    print(canonical(result))
    return result


INCOMPLETE = 'INCOMPLETE_SUPPORT'
UNRESOLVED = 'UNRESOLVED'


def conjunction(values):
    values = tuple(values)
    if 'FAIL' in values:
        return 'FAIL'
    if values and all(v == 'PASS' for v in values):
        return 'PASS'
    return INCOMPLETE if INCOMPLETE in values else UNRESOLVED


def upper(value, allowance, limit):
    if value is None or allowance is None:
        return INCOMPLETE
    return 'PASS' if value+allowance <= limit else 'FAIL' if value-allowance > limit else UNRESOLVED


METRICS = ('R_mg_g', 'B_mg_g', 'absB_mg_g', 'R_allowance_mg_g', 'B_allowance_mg_g', 'absB_allowance_mg_g')


def shot_metrics(rows, predictions):
    if (not rows or len(rows) != len(predictions)
            or any(not r['analyte_eligible'] or r.get('q') is None or r['mass_kg'] is None
                   or not p['numerical_qualified'] or p['five_cqa_mg_g'] is None
                   or not p['supported'] for r, p in zip(rows, predictions))):
        return None
    if any((r['shot'], r['fraction']) != (p['shot'], p['fraction']) for r, p in zip(rows, predictions)):
        raise ValueError('EXACT_PHYSICAL_SHOT_FRACTION_OUTCOME_JOIN_REQUIRED')
    mass = np.asarray([r['mass_kg'] for r in rows])
    error = np.asarray([p['five_cqa_mg_g']-1000*r['q'] for r, p in zip(rows, predictions)])
    allowance = np.asarray([1000*p['allowance_kg']/m for p, m in zip(predictions, mass)])
    if not np.isfinite([mass, error, allowance]).all() or np.any(mass <= 0) or np.any(allowance < 0):
        raise ValueError('FINITE_OBSERVATIONS_PREDICTIONS_AND_ALLOWANCES_REQUIRED')
    weights = mass/mass.sum()
    bias = float(weights @ error)
    return {'R_mg_g': float(np.sqrt(weights @ error**2)), 'B_mg_g': bias, 'absB_mg_g': abs(bias),
            'R_allowance_mg_g': float(np.sqrt(weights @ allowance**2)),
            'B_allowance_mg_g': float(weights @ allowance), 'absB_allowance_mg_g': float(weights @ allowance),
            'windows': len(rows), 'assayed_beverage_kg': float(mass.sum()),
            'observed_assayed_five_cqa_kg': float(sum(r['q']*r['mass_kg'] for r in rows)),
            'predicted_assayed_five_cqa_kg': float(sum(p['five_cqa_kg'] for p in predictions))}


def mean_metrics(values):
    if not values or any(v is None for v in values):
        return None
    return {k: float(np.mean([v[k] for v in values])) for k in METRICS}


def metric_bounds(value):
    if value is None:
        return None
    return {name: {'lower': max(0., value[name]-value[allowance]) if name != 'B_mg_g' else value[name]-value[allowance],
                   'upper': value[name]+value[allowance]}
            for name, allowance in [('R_mg_g', 'R_allowance_mg_g'), ('B_mg_g', 'B_allowance_mg_g'),
                                     ('absB_mg_g', 'absB_allowance_mg_g')]}


def panel_metrics(observed, predicted, conditions):
    pi = {(p['shot'], p['fraction']): p for p in predicted}
    if len(pi) != len(predicted):
        raise ValueError('DUPLICATE_PREDICTION_SLOT')
    ci, private = {}, []
    for condition in conditions:
        shots = sorted({r['shot'] for r in observed if r['condition'] == condition})
        if len(shots) != 3:
            raise ValueError('ORIGINAL_THREE_SHOT_DENOMINATOR_REQUIRED')
        full, diagnostic, coverage = [], [], []
        for shot in shots:
            rows = [r for r in observed if r['shot'] == shot]
            if len(rows) != 4 or {r['fraction'] for r in rows} != set(SUFFIX):
                raise ValueError('ORIGINAL_FOUR_FRACTION_DENOMINATOR_REQUIRED')
            predictions = [pi[(r['shot'], r['fraction'])] for r in rows]
            mask = [p['supported'] for p in predictions]
            restricted = shot_metrics([r for r, keep in zip(rows, mask) if keep],
                                      [p for p, keep in zip(predictions, mask) if keep])
            complete = shot_metrics(rows, predictions) if all(mask) else None
            full.append(complete); diagnostic.append(restricted)
            coverage.append({'intended': 4, 'supported': sum(mask),
                             'qualified': sum(p['numerical_qualified'] and keep for p, keep in zip(predictions, mask))})
            private.append({'shot': shot, 'condition': condition, 'full': complete, 'bounds': metric_bounds(complete),
                            'supported_subset_diagnostic': restricted, 'coverage': coverage[-1]})
        complete, restricted = mean_metrics(full), mean_metrics(diagnostic)
        lower_R = sum(max(0., v['R_mg_g']-v['R_allowance_mg_g']) for v in full if v)/3
        lower_A = sum(max(0., v['absB_mg_g']-v['absB_allowance_mg_g']) for v in full if v)/3
        adequacy = (conjunction([upper(complete['R_mg_g'], complete['R_allowance_mg_g'], .25),
                                upper(complete['absB_mg_g'], complete['absB_allowance_mg_g'], .125)])
                    if complete else 'FAIL' if lower_R > .25 or lower_A > .125 else INCOMPLETE)
        ci[condition] = {'full_scope_metrics': complete, 'bounds': metric_bounds(complete),
            'supported_subset_diagnostic': restricted, 'adequacy': adequacy, 'original_shots': 3,
            'original_slots': 12, 'complete_shots': sum(v is not None for v in full),
            'supported_slots': sum(v['supported'] for v in coverage), 'qualified_slots': sum(v['qualified'] for v in coverage),
            'complete_shot_full_condition_lower_bounds': {'R_mg_g': lower_R, 'absB_mg_g': lower_A}}
    balanced = mean_metrics([v['full_scope_metrics'] for v in ci.values()])
    return {'conditions': ci, 'original_conditions': len(conditions), 'original_shots': 3*len(conditions),
            'original_slots': 12*len(conditions), 'full_scope_metrics': balanced, 'bounds': metric_bounds(balanced),
            'supported_subset_diagnostic': mean_metrics([v['supported_subset_diagnostic'] for v in ci.values()]),
            'adequacy': conjunction(v['adequacy'] for v in ci.values()),
            'complete_scope': all(v['full_scope_metrics'] is not None for v in ci.values())}, private


def material_gain(candidate, control):
    if not candidate['complete_scope'] or not control['complete_scope']:
        return {'status': INCOMPLETE, 'reason': 'ALL_ORIGINAL_PRIMARY_SLOTS_REQUIRED'}
    c, k = candidate['full_scope_metrics'], control['full_scope_metrics']
    gain = k['R_mg_g']-c['R_mg_g']
    allowance = k['R_allowance_mg_g']+c['R_allowance_mg_g']
    relative = upper(c['R_mg_g']-.85*k['R_mg_g'], c['R_allowance_mg_g']+.85*k['R_allowance_mg_g'], 0.)
    wins = possible = 0
    for condition, entry in candidate['conditions'].items():
        a, b = entry['full_scope_metrics'], control['conditions'][condition]['full_scope_metrics']
        wins += a['R_mg_g']+a['R_allowance_mg_g'] < b['R_mg_g']-b['R_allowance_mg_g']
        possible += a['R_mg_g']-a['R_allowance_mg_g'] < b['R_mg_g']+b['R_allowance_mg_g']
    deterioration = c['absB_mg_g']-k['absB_mg_g']
    da = c['absB_allowance_mg_g']+k['absB_allowance_mg_g']
    components = {'absolute_gain': upper(-gain, allowance, -.05), 'relative_gain': relative,
                  'condition_wins': 'PASS' if wins >= 3 else 'FAIL' if possible < 3 else UNRESOLVED,
                  'abs_bias_deterioration': upper(deterioration, da, .025)}
    denominator_lower = k['R_mg_g']-k['R_allowance_mg_g']
    return {'status': conjunction(components.values()), 'components': components,
            'R_improvement_mg_g': gain, 'R_improvement_bounds': [gain-allowance, gain+allowance],
            'relative_improvement': gain/k['R_mg_g'] if k['R_mg_g'] > 0 else None,
            'relative_improvement_lower_bound': 1-(c['R_mg_g']+c['R_allowance_mg_g'])/denominator_lower if denominator_lower > 0 else None,
            'definite_condition_wins': int(wins), 'possible_condition_wins': int(possible),
            'mean_absB_deterioration_mg_g': deterioration, 'mean_absB_deterioration_bounds': [deterioration-da, deterioration+da]}


def disposition(adequacy, gain_status, complete):
    if not complete or any(v not in ('PASS', 'FAIL') for v in adequacy.values()):
        return 'NOT_ADJUDICATED', 'PRIMARY_SUPPORT_OR_ARM_ADEQUACY_UNRESOLVED'
    if adequacy['D0'] == 'PASS' and gain_status not in ('PASS', 'FAIL'):
        return 'NOT_ADJUDICATED', 'NUMERICAL_COMPLEXITY_COMPARISON_UNRESOLVED'
    if adequacy['D0'] == 'PASS' and gain_status == 'PASS':
        return 'FIVE_CQA_CONDITIONED_MODEL_EARNED', None
    if adequacy['E0'] == 'PASS':
        return 'SIMPLE_FIVE_CQA_EXPONENTIAL_ADEQUATE', None
    if adequacy['D0'] == 'PASS':
        return 'FIVE_CQA_D0_ADEQUATE_INCREMENT_NOT_ESTABLISHED', None
    return 'TESTED_FIVE_CQA_FAMILIES_INADEQUATE', None


def evaluate(observed, predictions):
    expected = {(r['shot'], r['fraction']) for r in observed}
    if len(observed) != 96 or len(expected) != 96 or set(predictions) != set(ARMS):
        raise ValueError('ALL_ORIGINAL_PRED_SLOTS_AND_BOTH_ARMS_REQUIRED')
    supports = []
    for ps in predictions.values():
        if len(ps) != 96 or {(p['shot'], p['fraction']) for p in ps} != expected:
            raise ValueError('MISSING_PREDICTIONS_CANNOT_DROP_DENOMINATORS')
        supports.append({(p['shot'], p['fraction']): (p['supported'], p['reason']) for p in ps})
    if supports[0] != supports[1]:
        raise ValueError('IDENTICAL_COMPARISON_SUPPORT_REQUIRED')
    panels, private = {}, {}
    for panel, conditions in PANELS.items():
        panels[panel], private[panel] = {}, {}
        for arm in ARMS:
            panels[panel][arm], private[panel][arm] = panel_metrics(observed, predictions[arm], conditions)
    primary = panels['primary']
    adequacy = {arm: primary[arm]['adequacy'] for arm in ARMS}
    gain = material_gain(primary['D0'], primary['E0'])
    label, blocker = disposition(adequacy, gain['status'], all(v['complete_scope'] for v in primary.values()))
    failures = {arm: dict(Counter(p['prediction_status'] for p in ps if not p['numerical_qualified']))
                for arm, ps in predictions.items()}
    extrapolation = {arm: {panel: {'shots': len({p['shot'] for p in ps if p['condition'] in conditions and p['feature_extrapolation']}),
                                  'features': sorted({f for p in ps if p['condition'] in conditions for f in p['feature_extrapolation']})}
                           for panel, conditions in PANELS.items()} for arm, ps in predictions.items()}
    return {'task': TASK, 'disposition': label, 'blocker': blocker, 'adequacy_by_arm': adequacy,
            'E0_to_D0_gain': gain, 'panels': panels, 'prediction_failures_and_unsupported': failures,
            'feature_extrapolation': extrapolation, 'claims': list(md.CLAIMS), 'rights': source.RIGHTS,
            'analytical_uncertainty': 'NOT_ESTABLISHED', 'physical_validation': 'NOT_ESTABLISHED',
            'scientific_score_passes': 1, 'native_ewp_runs': 0, 'merge_authorized': False,
            'production_adoption_authorized': False, 'successor_authorized': False}, private


def score_guard(out):
    if any((out / n).exists() for n in ('score_receipt.json', 'scores.json', 'score_completion.json',
                                      'observed_suffix.json', 'shot_results.json')):
        raise ValueError('DUPLICATE_SCORE_OR_PRESERVED_FAILED_ATTEMPT')


def verify_before_score(out, review):
    from . import pannusch_conditioned_mass_delivery as receipt
    score_guard(out)
    frozen = receipt.verify_before_score(out, review)
    approved = read(review)
    if (frozen['task'] != TASK or approved.get('task') != TASK or not approved.get('reviewer')
            or approved.get('unresolved_blocking_findings') != []
            or approved.get('future_chemistry_attached') is not False
            or approved.get('approval_scope') != 'ONE_PREDECLARED_FROZEN_PRED_SCORE_ONLY_NO_RETUNING'):
        raise ValueError('GENUINE_INDEPENDENT_EXACT_FREEZE_REVIEW_REQUIRED')
    verify_frozen(out)
    qualified_sources()
    return frozen


def score(out, review):
    check_task_clock()
    out = private_directory(out)
    verify_before_score(out, review)
    write(out / 'score_receipt.json', {'task': TASK, 'status': 'STARTED',
          'freeze_sha256': digest(out / 'freeze.json'), 'review_sha256': digest(Path(review)),
          'predictions_sha256': digest(out / 'predictions.json')})
    # Sole outcome join, after the independent approval and exclusive receipt.
    observed = read(out / 'queries.json')
    rows = source_geometry.all_rows()
    paths, _ = qualified_sources()
    index = analyte_index(original_rows(paths, rows, 'PREDICTION_2022_03'), 'PREDICTION_2022_03')
    for r in observed:
        original = index[r['shot'], r['fraction']]
        if not r['analyte_eligible'] or original['validity'] != 'VALID':
            raise ValueError('SOURCE_VALID_PRED_TARGET_REQUIRED')
        r['q'] = source_value(original, 'PREDICTION_2022_03')
    if original_reconciliation(paths, observed, rows) != 96:
        raise ValueError('ALL_PRED_FIVE_CQA_ORIGINALS_REQUIRED')
    result, private = evaluate(observed, read(out / 'predictions.json'))
    write(out / 'observed_suffix.json', observed)
    write(out / 'shot_results.json', private)
    write(out / 'scores.json', result)
    write(out / 'score_completion.json', {'task': TASK, 'status': 'COMPLETE', 'scientific_score_passes': 1,
          'files': {n: digest(out / n) for n in ('scores.json', 'shot_results.json', 'observed_suffix.json', 'score_receipt.json')},
          'freeze_sha256': digest(out / 'freeze.json'), 'review_sha256': digest(Path(review))})
    print(result['disposition'])


def report(out):
    out = private_directory(out)
    verify_frozen(out)
    completion = read(out / 'score_completion.json')
    if completion['task'] != TASK or completion['status'] != 'COMPLETE' or completion['scientific_score_passes'] != 1:
        raise ValueError('RETAINED_COMPLETED_SCORE_REQUIRED')
    if completion['freeze_sha256'] != digest(out / 'freeze.json'):
        raise ValueError('RETAINED_SCORE_FREEZE_MISMATCH')
    if set(completion['files']) != {'scores.json', 'shot_results.json', 'observed_suffix.json', 'score_receipt.json'}:
        raise ValueError('COMPLETE_RETAINED_SCORE_BINDING_REQUIRED')
    receipt = read(out / 'score_receipt.json')
    if (receipt['task'] != TASK or receipt['freeze_sha256'] != completion['freeze_sha256']
            or receipt['review_sha256'] != completion['review_sha256']
            or receipt['predictions_sha256'] != digest(out / 'predictions.json')):
        raise ValueError('RETAINED_SCORE_RECEIPT_BINDING_MISMATCH')
    for relative, sha in completion['files'].items():
        if digest(out / relative) != sha:
            raise ValueError('RETAINED_SCORE_ARTIFACT_DRIFT')
    result = read(out / 'scores.json')
    print(canonical(result))
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__, allow_abbrev=False)
    commands = {'prepare': prepare, 'develop': develop,
                'freeze': freeze, 'verify': verify, 'report': report}
    parser.add_argument('command', choices=tuple(commands)+('score',))
    parser.add_argument('--out', required=True, type=Path)
    parser.add_argument('--review', type=Path)
    args = parser.parse_args()
    if args.command == 'score':
        if args.review is None:
            parser.error('--review is required for score')
        score(args.out, args.review)
    else:
        if args.review is not None:
            parser.error('--review is accepted only for score')
        commands[args.command](args.out)


if __name__ == '__main__':
    main()
