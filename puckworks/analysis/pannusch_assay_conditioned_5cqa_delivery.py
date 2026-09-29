"""Bounded first-assay-conditioned 5-CQA research; originals and rows stay private."""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
import os
import platform
import subprocess
import time

import numpy as np
import scipy

from . import assay_conditioned_5cqa_delivery as md
from . import pannusch_conditional_5cqa_delivery as legacy
from . import pannusch_conditional_tail_delivery as source_geometry
from . import pannusch_mass_delivery as source

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT/'docs/analysis/sci_md_5cqa_assay_001'
TASK = 'SCI-MD-5CQA-ASSAY-001'
SPECIES = md.SPECIES
ARMS = ('E0', 'D0', 'A1', 'L1')
SUFFIX, PANELS = legacy.SUFFIX, legacy.PANELS
read, write, digest = legacy.read, legacy.write, legacy.digest
canonical, identity, private_directory = md.canonical, md.identity, legacy.private_directory
TASK_DEADLINE = datetime.fromisoformat('2026-09-29T18:22:32+00:00').timestamp()
INCOMPLETE, UNRESOLVED = legacy.INCOMPLETE, legacy.UNRESOLVED


def check_task_clock():
    if time.time() >= TASK_DEADLINE:
        raise RuntimeError('SIX_HOUR_TASK_DEADLINE')


def git(expression):
    return subprocess.check_output(['git', 'rev-parse', expression], cwd=ROOT, text=True).strip()


def contract():
    value = read(DOC/'INFORMATION_CONTRACT.json')
    if value['task'] != TASK or value['arms'] != list(ARMS) or value['new_heads'] != list(md.ARMS):
        raise ValueError('WRONG_FIRST_ASSAY_FIVE_CQA_CONTRACT')
    subprocess.check_call(['git', 'merge-base', '--is-ancestor', value['contract_commit'], 'HEAD'], cwd=ROOT)
    for relative, sha in value['files'].items():
        committed = subprocess.check_output(['git', 'show', value['contract_commit']+':'+relative], cwd=ROOT)
        if digest(ROOT/relative) != sha or source.hashlib.sha256(committed).hexdigest() != sha:
            raise ValueError('PROTOCOL_MUST_BE_COMMITTED_AND_IMMUTABLE_BEFORE_FITTING')
    return value


def qualified_sources():
    paths, history = legacy.qualified_sources()
    if history != read(DOC/'SOURCE_IDENTITIES.json'):
        raise ValueError('EXACT_EIGHT_ORIGINAL_TEN_REGISTER_BINDING_REQUIRED')
    for relative, sha in read(DOC/'DEPENDENCIES.json')['files'].items():
        if digest(ROOT/relative) != sha:
            raise ValueError('IMMUTABLE_PREDECESSOR_CHANGED:'+relative)
    return paths, history


def source_contract_error(reason, shot='', fraction=None):
    raise ValueError(f'SOURCE_CONTRACT_BLOCKED:{reason}:{shot}:fraction={fraction}')


def checked_source_row(row, campaign, shot, fraction):
    label, exp, rep = shot.split('-')
    expected_campaign = 'FIT_2021_12' if label == 'FIT' else 'PREDICTION_2022_03'
    if (campaign != expected_campaign or row.get('campaign_id') != campaign
            or row.get('shot_id') != shot or row.get('fraction_id') != str(fraction)
            or row.get('analyte') != SPECIES or row.get('source_id') != 'P24-MAT-'+label
            or row.get('condition_id') != label+'-C'+exp[1:]
            or row.get('physical_replicate_id') != rep[1:]):
        source_contract_error('WRONG_SOURCE_IDENTITY', shot, fraction)
    if label == 'FIT' and (row.get('fraction_basis') != 'MEASURED_MASS_G' or
            row.get('source_object_or_cell') != f'ExperimentalData({int(exp[1:])}).run({int(rep[1:])}) fraction {fraction}'):
        source_contract_error('WRONG_SOURCE_CELL_OR_MASS_BASIS', shot, fraction)
    if row.get('concentration_unit') != 'mg/g':
        source_contract_error('WRONG_ANALYTE_UNITS', shot, fraction)
    if row.get('validity') != 'VALID' or row.get('exclusion_reason'):
        source_contract_error('INVALID_REQUIRED_ASSAY', shot, fraction)


def first_projection(coords, rows):
    """Project only fraction 1 of exact 5CQA; q2/spills/other analytes stay out."""
    try:
        masses = legacy.mass_projection(coords)
    except ValueError as exc:
        source_contract_error('FIRST_VIAL_GEOMETRY:'+str(exc))
    chosen = [r for r in rows if r.get('analyte') == SPECIES and r.get('fraction_id') == '1']
    index = {r['shot_id']: r for r in chosen}
    if len(index) != len(chosen):
        source_contract_error('DUPLICATE_FIRST_ASSAY')
    if set(index) != set(masses):
        source_contract_error('MISSING_OR_EXTRA_FIRST_ASSAY:'+','.join(sorted(set(index)^set(masses))))
    result, projected = {}, []
    for shot, values in masses.items():
        row = index[shot]
        campaign = 'FIT_2021_12' if shot.startswith('FIT-') else 'PREDICTION_2022_03'
        checked_source_row(row, campaign, shot, 1)
        field = 'concentration_value' if campaign == 'FIT_2021_12' else 'measured_concentration'
        try:
            q = md.first_assay_from_mg_g(float(row[field]))
            grams = md.number(float(row['fraction_liquid_g_or_ml']))
        except (ValueError, TypeError, KeyError) as exc:
            source_contract_error('MISSING_OR_INVALID_FIRST_ASSAY:'+str(exc), shot, 1)
        if abs(grams/1000-values[0]) > 5.001e-12:
            source_contract_error('FIRST_VIAL_MASS_EXPORT_MISMATCH', shot, 1)
        result[shot] = values+[q]
        projected.append({'campaign': campaign, 'shot': shot, 'fraction': 1,
                          'species': SPECIES, 'source_id': row['source_id'],
                          'mass_kg': values[0], 'q': q})
    return result, projected


def formula_cell(formulas, cached, label, experiment, replicate, fraction):
    """Check exactly the permitted fraction's campaign calibration and dilution."""
    import re
    sheet = f'{3*experiment-2}-{3*experiment}'
    fs, cs = formulas[sheet], cached[sheet]
    k = legacy.ASSAY_FRACTIONS.index(fraction)
    row = 10+6*(replicate-1)+k
    if fs[8][24].value != 'Chlorogenic acid':
        source_contract_error('ORIGINAL_ANALYTE_HEADER', f'{label}-E{experiment:02d}-R{replicate}', fraction)
    sign, offset, slope = ('-', 78.923, 24.513) if label == 'FIT' else ('+', 88.067, 26.383)
    formula = fs[row-1][24].value
    match = re.match(r'=\(I(\d+)', str(formula))
    if not match:
        source_contract_error('UNRESOLVED_CALIBRATION_FORMULA', sheet, fraction)
    origin = int(match[1])
    if formula != f'=(I{origin}{sign}{offset})/({slope}*1000)*$C{origin}*$B{origin}':
        source_contract_error('CALIBRATION_DILUTION_FORMULA_CHANGED', sheet, fraction)
    mass, dilution, area = [md.number(cs[origin-1][col].value) for col in (1, 2, 8)]
    expected = (area+(-offset if label == 'FIT' else offset))/(slope*1000)*dilution*mass
    actual = md.number(cs[row-1][24].value)
    if mass <= 0 or dilution <= 0 or expected < 0 or abs(expected-actual) > 1e-12:
        source_contract_error('CALIBRATION_RECONCILIATION', sheet, fraction)
    return actual


def source_books(path):
    import openpyxl
    import re
    result = []
    for cached in (False, True):
        book = openpyxl.load_workbook(path, read_only=True, data_only=cached)
        result.append({n: list(book[n].iter_rows()) for n in book.sheetnames
                       if re.fullmatch(r'\d+-\d+', n)})
        book.close()
    return result


def original_first_assays(paths, early, projected):
    """Only first-fraction MAT/HPLC cells cross this input authority boundary."""
    objects, books, result = {}, {}, {}
    counts = {'FIT': 0, 'PRED': 0}
    for row in projected:
        shot = row['shot']; label, exp, rep = shot.split('-')
        if row['fraction'] != 1 or row['species'] != SPECIES:
            source_contract_error('FIRST_ASSAY_ROLE_REQUIRED', shot, row['fraction'])
        if label not in objects:
            objects[label] = source.mat(paths['P24-MAT-'+label])['ExperimentalData']
            books[label] = source_books(paths['P24-HPLC-'+label])
        e, j = int(exp[1:]), int(rep[1:])
        run = objects[label][e-1].run[j-1]
        original = md.number(float(run.cAlcaloids[0, 2]))
        mass_g = md.number(float(run.mE[0]))
        assay_mass_mg = formula_cell(*books[label], label, e, j, 1)
        q = md.first_assay_from_mg_g(original)
        if (mass_g <= 0 or abs(original-1000*row['q']) > 5.001e-9
                or abs(mass_g/1000-row['mass_kg']) > 5.001e-12
                or abs(assay_mass_mg/mass_g-original) > 1e-12):
            source_contract_error('ORIGINAL_FIRST_ASSAY_RECONCILIATION', shot, 1)
        result[shot] = list(early[shot][:2])+[q]
        counts[label] += 1
    if counts != {'FIT': 45, 'PRED': 24}:
        source_contract_error('ALL_SIXTY_NINE_FIRST_ASSAYS_REQUIRED')
    return result, {'first_fraction_reconciliations': counts,
                    'first_fraction_HPLC_formulas': dict(counts), 'fraction_2_assays_projected': 0,
                    'public_rounding_is_analytical_uncertainty': False}


def suffix_formula_check(paths, label):
    """FIT-only before fit; PRED-only after approved score receipt."""
    if label not in ('FIT', 'PRED'):
        raise ValueError('DECLARED_CAMPAIGN_REQUIRED')
    formulas, cached = source_books(paths['P24-HPLC-'+label])
    count = 0
    for e in range(1, (15 if label == 'FIT' else 8)+1):
        for rep in (1, 2, 3):
            for fraction in SUFFIX:
                formula_cell(formulas, cached, label, e, rep, fraction)
                count += 1
    return {'qualified_suffix_formula_cells': count, 'status': 'PASS'}


def pred_queries(coords, rows):
    """PRED suffix metadata and mass geometry only; never read a concentration."""
    selected = [r for r in rows if r['campaign_id'] == 'PREDICTION_2022_03'
                and r['analyte'] == SPECIES and int(r['fraction_id']) in SUFFIX]
    index = {(r['shot_id'], int(r['fraction_id'])): r for r in selected}
    if len(index) != len(selected):
        source_contract_error('DUPLICATE_PRED_SUFFIX_IDENTITY')
    slots = []
    for c in coords:
        if c['campaign'] != 'PREDICTION_2022_03' or c['fraction'] not in SUFFIX:
            continue
        row = index.get((c['shot'], c['fraction']))
        if row is None:
            source_contract_error('MISSING_PRED_SUFFIX_IDENTITY', c['shot'], c['fraction'])
        checked_source_row(row, c['campaign'], c['shot'], c['fraction'])
        _, exp, rep = c['shot'].split('-')
        item = {k: c[k] for k in ('campaign', 'condition', 'shot', 'fraction', 'b0', 'b1',
                                  'mass_kg', 'coordinate_status')}
        item.update(analyte=SPECIES, source_id=row['source_id'],
                    source_field=f'ExperimentalData({int(exp[1:])}).run({int(rep[1:])}).cAlcaloids({legacy.ASSAY_FRACTIONS.index(c["fraction"])+1},3)',
                    analyte_eligible=True, source_reason='', validity_reason='VALID_SOURCE_FIVE_CQA')
        slots.append(item)
    if len(slots) != len(index):
        source_contract_error('EXTRA_PRED_SUFFIX_IDENTITY')
    return slots


def prepare(out):
    check_task_clock()
    authority = contract()
    paths, history = qualified_sources()
    out = private_directory(out, create=True); out.chmod(0o700)
    coords, rows = source_geometry.coordinates(paths), source_geometry.all_rows()
    early, projected = first_projection(coords, rows)
    early, first_audit = original_first_assays(paths, early, projected)
    fit_rows = [r for r in rows if r['campaign_id'] == 'FIT_2021_12']
    for r in fit_rows:
        if r['analyte'] == SPECIES and int(r['fraction_id']) in SUFFIX:
            checked_source_row(r, r['campaign_id'], r['shot_id'], int(r['fraction_id']))
    qualified_fit = legacy.original_rows(paths, fit_rows, 'FIT_2021_12')
    projections = build_projections(coords,
        qualified_fit+[r for r in rows if r['campaign_id'] == 'PREDICTION_2022_03'],
        source_geometry.condition_groups(), early=early)
    counts = projections['source_counts']
    expected = read(DOC/'DATA_AVAILABILITY_PREFLIGHT.json')['counts']
    if counts != expected:
        source_contract_error('QUALIFIED_SOURCE_COHORT_OR_SUPPORT_CHANGED')
    reconciled = legacy.original_reconciliation(paths, projections['fit_slots'], fit_rows)
    if reconciled != 180:
        source_contract_error('ALL_FIT_SUFFIX_ORIGINALS_REQUIRED')
    counts.update(first_audit, FIT_suffix_original_checks=reconciled,
                  FIT_suffix_formulas=suffix_formula_check(paths, 'FIT'))
    projections.update(source=dict(history, rights=source.RIGHTS), information_contract=authority)
    for name, value in projections.items():
        write(out/(name+'.json'), value)
    write(out/'prepared_hashes.json', {str(p.relative_to(out)): digest(p) for p in out.rglob('*.json')})
    write(out/'preparation_runtime.json', source_geometry.imported_files())
    print(canonical(counts))


def develop(out):
    from . import assay_conditioned_5cqa_training as train
    check_task_clock(); out = prepared(out)
    if any(os.environ.get(k) != '1' for k in ('OPENBLAS_NUM_THREADS', 'OMP_NUM_THREADS', 'MKL_NUM_THREADS')):
        raise ValueError('ONE_BLAS_THREAD_REQUIRED')
    dest = out/'development'; dest.mkdir()
    budget = train.Budget(out/'starts', evidence_root=out.parent)
    try:
        records, source_id = read(out/'training.json'), source_identity(out)
        specs = read(out/'development_support.json')
        shots = train.project_arm(records)
        candidates = []
        for lam in train.LAMBDAS:
            candidate = {'lambda': lam, 'status': 'SELECTABLE', 'folds': []}
            for spec in specs:
                group = spec['group']
                training, held = train.split_design(shots, group)
                support = set(map(tuple, spec['support']))
                expected = {(s.shot, w.fraction) for s in held for w in s.windows if w.end_kg <= spec['domain_kg']}
                if ([s.shot for s in training] != spec['training_shots']
                        or [s.shot for s in held] != spec['held_shots'] or support != expected
                        or max(w.end_kg for s in training for w in s.windows) != spec['domain_kg']):
                    raise ValueError('WHOLE_ORIGINAL_DESIGN_FOLD_OR_SUPPORT_DRIFT')
                name = f'L1.{lam}.{group}'
                model, audit = train.fit(training, lam, budget, name, source_id, source.RIGHTS)
                write(dest/(name+'.audit.json'), audit)
                fold = {'group': group, 'fit_status': audit['status'],
                        'domain_kg': spec['domain_kg'], 'training_shots': spec['training_shots'],
                        'held_shots': spec['held_shots'], 'support': spec['support']}
                try:
                    if model is None:
                        raise ValueError('FIT_NOT_QUALIFIED:'+audit['status'])
                    model.save(dest/(name+'.model.json'))
                    fold.update(train.held_metrics(model, held, support), model_sha256=model.sha256)
                except ValueError as exc:
                    candidate['status'] = 'NONSELECTABLE'; fold['failure'] = str(exc)
                candidate['folds'].append(fold)
                write(dest/(name+'.fold.json'), fold)
                print(name, fold['fit_status'], flush=True)
            for target, key in [('balanced_R_mg_g', 'R_mg_g'), ('R_allowance_mg_g', 'R_allowance_mg_g')]:
                candidate[target] = float(np.mean([f[key] for f in candidate['folds']])) if candidate['status'] == 'SELECTABLE' else None
            candidates.append(candidate)
            write(dest/f'L1.{lam}.candidate.json', candidate)
        lam = train.select_lambda(candidates)
        summary = {'task': TASK, 'scope': 'FIT_ONLY_MODEL_SELECTION_NOT_EXTERNAL_VALIDATION',
                   'arms': {'L1': {'candidates': candidates, 'selected_lambda': lam}},
                   'A1_development_CV': 'NOT_PERFORMED_DECAY_USES_ALL_FIT'}
        write(dest/'development.json', summary)
        model_dir = out/'models'; model_dir.mkdir()
        model, audit = train.fit(shots, lam, budget, 'L1.final', source_id, source.RIGHTS, scope='FINAL_FIT')
        write(model_dir/'L1.audit.json', audit)
        if model is None:
            raise ValueError('SUPPORT_OR_NUMERICAL_RESULT_NOT_ADJUDICATED:FINAL_FIT')
        model.save(model_dir/'L1.json')
        support = {(s.shot, w.fraction) for s in shots for w in s.windows}
        write(model_dir/'L1.qualification.json', train.held_metrics(model, shots, support))
        # Feature summaries only: A1 has no new fitted parameter or CV score.
        info = {'species': SPECIES, 'scope': 'FIXED_E0_DECAY_QUERY_FIRST_ASSAY_NORMALIZATION',
                'E0_semantic_sha256': md.E0_SHA256, 'source_identity': source_id,
                'feature_summary_shots': [s.shot for s in shots],
                'feature_summary_groups': sorted({s.group for s in shots}), 'global_fits': 0}
        a1 = md.Model('A1', (), model.means, model.minima, model.maxima, model.domain_kg,
                      None, canonical(info), source.RIGHTS)
        a1.save(model_dir/'A1.json')
        write(model_dir/'A1.qualification.json', train.held_metrics(a1, shots, support))
        for arm in legacy.ARMS:
            src = ROOT/'docs/analysis/sci_md_5cqa_delivery_001/models'/f'{arm}.json'
            with (model_dir/(arm+'.json')).open('xb') as stream:
                stream.write(src.read_bytes())
        if legacy.md.Model.load(model_dir/'E0.json').sha256 != md.E0_SHA256:
            raise ValueError('EXACT_E0_DECAY_AUTHORITY_REQUIRED')
        write(out/'fitting_completion.json', {'status': 'COMPLETE',
              'fitting_stage_wall_seconds': time.time()-budget.clock['started_unix']})
    finally:
        budget.worker_lock.close()


def decide(adequacy, comparisons, complete):
    adequate = [a for a in md.ARMS if adequacy[a] == 'PASS']
    least = adequate[0] if adequate else None
    earned = adequacy['L1'] == 'PASS' and all(comparisons[k]['status'] == 'PASS'
                                           for k in ('L1_vs_D0', 'L1_vs_A1'))
    blocker = None
    if not complete or any(adequacy[a] not in ('PASS', 'FAIL') for a in md.ARMS):
        label = 'SUPPORT_OR_NUMERICAL_RESULT_NOT_ADJUDICATED'
        blocker = 'PRIMARY_SUPPORT_NUMERICS_OR_ADEQUACY_THRESHOLD_OVERLAP'
    elif adequate and any(c['status'] not in ('PASS', 'FAIL') for c in comparisons.values()):
        label = 'SUPPORT_OR_NUMERICAL_RESULT_NOT_ADJUDICATED'
        blocker = 'MATERIAL_INCREMENT_NUMERICAL_AMBIGUITY'
    elif earned:
        label = 'L1_ADEQUATE_AND_EARNED'
    elif least == 'A1':
        label = 'A1_ADEQUATE_LEARNED_INCREMENT_NOT_ESTABLISHED'
    elif least:
        label = 'ADEQUATE_INCREMENT_NOT_ESTABLISHED'
    else:
        label = 'TESTED_FIRST_ASSAY_FIVE_CQA_FAMILIES_INADEQUATE'
    return {'disposition': label, 'blocker': blocker, 'least_complex_adequate_new_arm': least,
            'adequate_new_arms': adequate, 'learned_route_earned': earned,
            'earned_replacement_arm': 'L1' if earned else None}


def verify_historical_results(result):
    """Compare retained aggregates only; never invoke the predecessor scorer."""
    retained = read(legacy.DOC/'RESULTS.json')
    maximum = 0.
    for panel in PANELS:
        for arm in legacy.ARMS:
            new, old = result['panels'][panel][arm], retained['panels'][panel][arm]
            if (new['adequacy'] != old['adequacy'] or new['complete_scope'] != old['complete_scope']
                    or new['original_slots'] != old['original_slots']):
                raise ValueError('HISTORICAL_CONTROL_DISPOSITION_DRIFT')
            for condition in PANELS[panel]:
                a, b = new['conditions'][condition], old['conditions'][condition]
                for key in ('supported_slots', 'qualified_slots', 'original_slots', 'adequacy'):
                    if a[key] != b[key]:
                        raise ValueError('HISTORICAL_CONTROL_SUPPORT_OR_DECISION_DRIFT')
                for role in ('full_scope_metrics', 'supported_subset_diagnostic'):
                    x, y = a[role], b[role]
                    if (x is None) != (y is None):
                        raise ValueError('HISTORICAL_CONTROL_METRIC_SCOPE_DRIFT')
                    if x is None:
                        continue
                    for key, allowance in [('R_mg_g', 'R_allowance_mg_g'), ('B_mg_g', 'B_allowance_mg_g'),
                                           ('absB_mg_g', 'absB_allowance_mg_g')]:
                        delta = abs(x[key]-y[key]); maximum = max(maximum, delta)
                        if delta > x[allowance]+y[allowance]+1e-12:
                            raise ValueError('HISTORICAL_CONTROL_RESULT_DRIFT')
    return {'status': 'PASS', 'maximum_difference_mg_g': maximum,
            'allowance': 'COMBINED_RETAINED_NUMERICAL_BOUNDS_PLUS_1e-12_mg_g',
            'predecessor_scoring_calls': 0}


def build_projections(coords, rows, groups, *, early=None):
    projected, _ = first_projection(coords, rows)
    if early is None:
        early = projected
    if set(early) != set(projected) or any(early[s][:2] != projected[s][:2] for s in early):
        source_contract_error('EXACT_FIRST_ASSAY_MASS_JOIN_REQUIRED')
    fit_slots = legacy.target_slots(coords, rows, campaign='FIT_2021_12', include_values=True)
    queries = pred_queries(coords, rows)
    training = []
    for shot in sorted(s for s in early if s.startswith('FIT-')):
        slots = [s for s in fit_slots if s['shot'] == shot]
        windows = []
        for s in slots:
            reason = s['source_reason'] or (s['coordinate_status'] if s['coordinate_status'] != 'QUALIFIED' else '')
            s.update(supported=not reason, reason=reason)
            if not reason:
                anchor = type('Anchor', (), {'b_anchor': sum(early[shot][:2])})()
                start, allowance = source_geometry.source_query_start(s, anchor)
                s.update(integration_start_kg=start, coordinate_allowance_kg=allowance)
                windows.append({'fraction': s['fraction'], 'start_kg': start, 'end_kg': s['b1'],
                                'q': s['q'], 'mass_kg': s['mass_kg'], 'coordinate_allowance_kg': allowance})
        if not windows:
            raise ValueError('ORIGINAL_FIT_SHOT_WITHOUT_SUPERVISION')
        training.append({'shot': shot, 'group': groups[slots[0]['condition']],
                         'early_values': early[shot], 'windows': windows})
    domain = max(w['end_kg'] for s in training for w in s['windows'])
    for q in queries:
        reason = q['source_reason'] or (q['coordinate_status'] if q['coordinate_status'] != 'QUALIFIED' else
                 'OUTSIDE_COMMON_TRAINING_MASS_DOMAIN' if q['b1'] > domain else '')
        q.update(supported=not reason, reason=reason)
    originals = sorted({s['group'] for s in training})
    if len(training) != 45 or len(originals) != 15 or len(fit_slots) != 180 or len(queries) != 96:
        raise ValueError('ORIGINAL_SOURCE_DENOMINATORS_REQUIRED')
    for group in originals:
        if sum(len(s['windows']) >= 3 for s in training if s['group'] == group) < 2:
            raise ValueError('TWO_PHYSICAL_SHOTS_WITH_THREE_TARGET_WINDOWS_REQUIRED')
    primary = [q for q in queries if q['condition'] in PANELS['primary']]
    if len(primary) != 48 or len({q['shot'] for q in primary}) != 12 or not all(q['supported'] for q in primary):
        raise ValueError('ALL_FORTY_EIGHT_PRIMARY_WINDOWS_REQUIRED')
    folds = []
    for group in originals:
        retained = [s for s in training if s['group'] != group]
        held = [s for s in training if s['group'] == group]
        bound = max(w['end_kg'] for s in retained for w in s['windows'])
        support = [(s['shot'], w['fraction']) for s in held for w in s['windows'] if w['end_kg'] <= bound]
        held_ids = {s['shot'] for s in held}
        folds.append({'group': group, 'domain_kg': bound, 'training_shots': [s['shot'] for s in retained],
                      'held_shots': [s['shot'] for s in held], 'support': support,
                      'slots': [{'shot': s['shot'], 'fraction': s['fraction'],
                                 'supported': (s['shot'], s['fraction']) in support,
                                 'reason': s['reason'] or ('OUTSIDE_FOLD_DOMAIN' if (s['shot'], s['fraction']) not in support else '')}
                                for s in fit_slots if s['shot'] in held_ids]})
    counts = {'fit_physical_shots': 45, 'fit_original_designs': 15, 'fit_intended_slots': 180,
              'fit_supported_slots': sum(s['supported'] for s in fit_slots),
              'fit_reasons': dict(Counter(s['reason'] for s in fit_slots if s['reason'])),
              'common_domain_kg': [0., domain], 'PRED_suffix_values_projected': 0, 'PRED_first_assay_values_projected': 24,
              'panels': {p: {'original_shots': 3*len(cs), 'intended_slots': 12*len(cs),
                             'supported_slots': sum(q['supported'] for q in queries if q['condition'] in cs)}
                         for p, cs in PANELS.items()}}
    clean_coords = [{k: c[k] for k in ('campaign', 'condition', 'shot', 'fraction', 'b0', 'b1',
                                     'mass_kg', 'coordinate_status')} for c in coords]
    return {'training': training, 'fit_slots': fit_slots, 'queries': queries,
            'early_inputs': {k: v for k, v in early.items() if k.startswith('PRED-')},
            'development_support': folds, 'source_counts': counts, 'coordinates': clean_coords}


def prepared(out):
    out = private_directory(out)
    contract(); qualified_sources()
    if any((out/n).exists() for n in ('freeze.json', 'score_receipt.json')):
        raise ValueError('FROZEN_RUN_IMMUTABLE')
    for relative, sha in read(out/'prepared_hashes.json').items():
        if digest(out/relative) != sha:
            raise ValueError('PREPARED_PROJECTION_DRIFT')
    for relative, sha in read(out/'preparation_runtime.json').items():
        if digest(ROOT/relative) != sha:
            raise ValueError('PREPARATION_RUNTIME_DRIFT')
    return out


def source_identity(out):
    value = read(out/'source.json')
    return identity({'source_files': {k: v for k, v in value['source_files'].items() if k.endswith('-FIT')},
                     'fit_register': value['registers']['fit_fraction_replicates.csv']})


def load_models(out):
    return {a: (legacy.md.Model if a in legacy.ARMS else md.Model).load(out/'models'/(a+'.json')) for a in ARMS}


def predict_records(models, early, queries):
    if set(models) != set(ARMS):
        raise ValueError('COMPLETE_FOUR_ARM_MATRIX_REQUIRED')
    predictions, states = legacy.predict_records({a: models[a] for a in legacy.ARMS},
                                                  {k: v[:2] for k, v in early.items()}, queries)
    for arm in md.ARMS:
        results = []
        for q in queries:
            # The legacy operator above already checked exact query keys/species.
            result = dict(q, prediction_status=q['reason'] or 'NOT_PREDICTED', numerical_qualified=False,
                          five_cqa_kg=None, five_cqa_mg=None, five_cqa_mg_g=None,
                          allowance_kg=None, feature_extrapolation=[])
            if q['supported']:
                try:
                    state = models[arm].condition(md.EarlyInput(*early[q['shot']], input_class='SOURCE_EARLY_INPUT'))
                    start, coord_allowance = source_geometry.source_query_start(q, state)
                    p = state.predict_intervals([start], [q['b1']])[0]
                    allowance = p.allowance_kg+coord_allowance
                    result.update(asdict(p), integration_start_kg=start, coordinate_allowance_kg=coord_allowance,
                                  allowance_kg=allowance, numerical_qualified=p.numerical_qualified and allowance <= 1e-9,
                                  prediction_status='QUALIFIED' if p.numerical_qualified and allowance <= 1e-9 else 'NUMERICALLY_UNQUALIFIED')
                    states[arm+':'+q['shot']] = state.to_dict()
                except (ValueError, FloatingPointError) as exc:
                    result['prediction_status'] = 'PREDICTION_FAILED:'+str(exc)
            results.append(result)
        predictions[arm] = results
    return predictions, states


def execution_account(out):
    starts = list((out/'starts').glob('*.start.json'))
    logs = [read(p) for p in (out/'starts').glob('*.end.json')]
    return {'task': TASK, 'iterative_starts': len(starts), 'completed_starts': len(logs),
            'starts_by_arm': {a: len(list((out/'starts').glob(a+'.*.start.json'))) for a in ('L1',)},
            'A1_global_fits': 0,
            'parent_fits': 0, 'E0_D0_fits': 0,
            'fitting_stage_wall_seconds': read(out/'fitting_completion.json')['fitting_stage_wall_seconds'],
            'actual_residual_calls': sum(r['actual_residual_calls'] for r in logs),
            'numerical_jacobian_residual_calls': sum(r['numerical_jacobian_residual_calls'] for r in logs),
            'max_calls_per_start': max((r['actual_residual_calls'] for r in logs), default=0),
            'optimizer_wall_seconds': sum(r['elapsed_seconds'] for r in logs),
            'failed_starts': sum(r['status'] != 'CONVERGED' for r in logs),
            'boundary_starts': sum(bool(r['boundary_indices']) for r in logs),
            'peak_rss_kib': max((r['peak_rss_kib'] for r in logs), default=0),
            'workers': 1, 'blas_threads': 1, 'native_ewp_runs': 0, 'scientific_score_passes': 0,
            'python': platform.python_version(), 'numpy': np.__version__, 'scipy': scipy.__version__,
            'imported_first_party_modules': source_geometry.imported_files()}


def freeze(out):
    check_task_clock(); out = prepared(out)
    if subprocess.check_output(['git', 'status', '--porcelain'], cwd=ROOT, text=True).strip():
        raise ValueError('COMMITTED_CLEAN_PRODUCER_REQUIRED')
    execution = execution_account(out)
    if (execution['iterative_starts'] != 183 or execution['completed_starts'] != 183
            or execution['starts_by_arm'] != {'L1': 183}
            or execution['max_calls_per_start'] > 8000 or execution['fitting_stage_wall_seconds'] >= 3600):
        raise ValueError('EXECUTION_ACCOUNTING_NOT_QUALIFIED')
    predictions, states = predict_records(load_models(out), read(out/'early_inputs.json'), read(out/'queries.json'))
    write(out/'predictions.json', predictions); write(out/'states.json', states)
    write(out/'legacy_predictions.json', {a: predictions[a] for a in legacy.ARMS})
    expected = read(ROOT/'docs/analysis/sci_md_5cqa_delivery_001/FREEZE.json')['artifacts']['predictions.json']
    if digest(out/'legacy_predictions.json') != expected:
        raise ValueError('FROZEN_E0_D0_REPLAY_MISMATCH')
    for ps in predictions.values():
        primary = [p for p in ps if p['condition'] in PANELS['primary']]
        if len(ps) != 96 or len(primary) != 48 or not all(p['supported'] and p['numerical_qualified'] for p in primary):
            raise ValueError('ALL_ORIGINAL_QUALIFIED_PRIMARY_WINDOWS_REQUIRED')
        if any(not p['numerical_qualified'] for p in ps if p['supported']):
            raise ValueError('NUMERICALLY_UNQUALIFIED_SUPPORTED_PREDICTION')
    execution.update(max_allowance_kg=max(p['allowance_kg'] for ps in predictions.values() for p in ps if p['allowance_kg'] is not None),
                     prediction_slots=384, legacy_prediction_sha256=expected, legacy_replay='EXACT_RETAINED_BYTES')
    write(out/'execution.json', execution)
    files = set(execution['imported_first_party_modules']) | set(read(out/'preparation_runtime.json'))
    files |= set(read(DOC/'DEPENDENCIES.json')['files'])
    files |= {str(p.relative_to(ROOT)) for p in DOC.rglob('*') if p.is_file()}
    files |= {str(p.relative_to(ROOT)) for p in (ROOT/'tests').glob('test_*assay*5cqa*.py')}
    files |= {'puckworks/analysis/assay_conditioned_5cqa_training.py'}
    write(out/'freeze.json', {'task': TASK, 'producer_commit': git('HEAD'), 'producer_tree': git('HEAD^{tree}'),
        'code_and_protocol': {p: digest(ROOT/p) for p in sorted(files)},
        'artifacts': {str(p.relative_to(out)): digest(p) for p in sorted(out.rglob('*.json'))},
        'arms': list(ARMS), 'panels': PANELS, 'outcomes_attached': False,
        'review_status': 'READY_FOR_INDEPENDENT_REVIEW', 'scoring_policy': 'ONE_INDEPENDENTLY_APPROVED_SCORE_NO_RETUNING'})
    print('READY_FOR_INDEPENDENT_REVIEW', digest(out/'freeze.json'))


def verify_frozen(out):
    frozen = read(out / 'freeze.json')
    if frozen['task'] != TASK or frozen['arms'] != list(ARMS) or frozen['outcomes_attached'] is not False:
        raise ValueError('WRONG_FROZEN_EXPERIMENT')
    for relative, sha in frozen['code_and_protocol'].items():
        if digest(ROOT / relative) != sha:
            raise ValueError('FROZEN_CODE_OR_PROTOCOL_DRIFT:' + relative)
        committed = subprocess.check_output(['git', 'show', frozen['producer_commit']+':'+relative], cwd=ROOT)
        if source.hashlib.sha256(committed).hexdigest() != sha:
            raise ValueError('FROZEN_PRODUCER_COMMIT_BINDING_MISMATCH')
    if git(frozen['producer_commit']+'^{tree}') != frozen['producer_tree']:
        raise ValueError('FROZEN_PRODUCER_TREE_MISMATCH')
    for relative, sha in frozen['artifacts'].items():
        if digest(out / relative) != sha:
            raise ValueError('FROZEN_ARTIFACT_DRIFT:' + relative)
    source.verify_registers(read(out / 'source.json'))
    return frozen


def verify(out):
    out = private_directory(out)
    frozen = verify_frozen(out)
    models = load_models(out)
    predictions, states = predict_records(models, read(out / 'early_inputs.json'), read(out / 'queries.json'))
    if canonical(predictions) != canonical(read(out / 'predictions.json')) or canonical(states) != canonical(read(out / 'states.json')):
        raise ValueError('SAVED_MODEL_INFERENCE_NOT_DETERMINISTIC')
    result = {'task': TASK, 'status': 'PASS', 'reviewed_producer': frozen['producer_commit'],
              'prediction_slots': sum(map(len, predictions.values())), 'prediction_max_difference': 0.,
              'optimizer_calls': 0, 'outcome_joins': 0, 'score_calls': 0, 'writes': 0}
    print(canonical(result))
    return result


def score_guard(out):
    if any((out / n).exists() for n in ('score_receipt.json', 'scores.json', 'score_completion.json',
                                      'observed_suffix.json', 'shot_results.json')):
        raise ValueError('DUPLICATE_SCORE_OR_PRESERVED_FAILED_ATTEMPT')


def verify_before_score(out, review):
    from . import pannusch_conditioned_mass_delivery as receipt
    score_guard(out)
    frozen = receipt.verify_before_score(out, review)
    approved = read(review)
    if (frozen['task'] != TASK or approved.get('task') != TASK or not approved.get('reviewer')
            or approved.get('unresolved_blocking_findings') != []
            or approved.get('future_chemistry_attached') is not False
            or approved.get('approval_scope') != 'ONE_PREDECLARED_FROZEN_PRED_SCORE_ONLY_NO_RETUNING'):
        raise ValueError('GENUINE_INDEPENDENT_EXACT_FREEZE_REVIEW_REQUIRED')
    verify_frozen(out)
    qualified_sources()
    return frozen


def score(out, review):
    check_task_clock()
    out = private_directory(out)
    verify_before_score(out, review)
    write(out / 'score_receipt.json', {'task': TASK, 'status': 'STARTED',
          'freeze_sha256': digest(out / 'freeze.json'), 'review_sha256': digest(Path(review)),
          'predictions_sha256': digest(out / 'predictions.json')})
    # Sole outcome join, after the independent approval and exclusive receipt.
    observed = read(out / 'queries.json')
    rows = source_geometry.all_rows()
    paths, _ = qualified_sources()
    index = legacy.analyte_index(legacy.original_rows(paths, rows, 'PREDICTION_2022_03'), 'PREDICTION_2022_03')
    for r in observed:
        original = index[r['shot'], r['fraction']]
        if not r['analyte_eligible'] or original['validity'] != 'VALID':
            raise ValueError('SOURCE_VALID_PRED_TARGET_REQUIRED')
        r['q'] = legacy.source_value(original, 'PREDICTION_2022_03')
    if legacy.original_reconciliation(paths, observed, rows) != 96:
        raise ValueError('ALL_PRED_FIVE_CQA_ORIGINALS_REQUIRED')
    suffix_formula_check(paths, 'PRED')
    result, private = evaluate(observed, read(out / 'predictions.json'))
    result['historical_control_agreement'] = verify_historical_results(result)
    write(out / 'observed_suffix.json', observed)
    write(out / 'shot_results.json', private)
    write(out / 'scores.json', result)
    write(out / 'score_completion.json', {'task': TASK, 'status': 'COMPLETE', 'scientific_score_passes': 1,
          'files': {n: digest(out / n) for n in ('scores.json', 'shot_results.json', 'observed_suffix.json', 'score_receipt.json')},
          'freeze_sha256': digest(out / 'freeze.json'), 'review_sha256': digest(Path(review))})
    print(result['disposition'])


def report(out):
    out = private_directory(out)
    verify_frozen(out)
    completion = read(out / 'score_completion.json')
    if completion['task'] != TASK or completion['status'] != 'COMPLETE' or completion['scientific_score_passes'] != 1:
        raise ValueError('RETAINED_COMPLETED_SCORE_REQUIRED')
    if completion['freeze_sha256'] != digest(out / 'freeze.json'):
        raise ValueError('RETAINED_SCORE_FREEZE_MISMATCH')
    if set(completion['files']) != {'scores.json', 'shot_results.json', 'observed_suffix.json', 'score_receipt.json'}:
        raise ValueError('COMPLETE_RETAINED_SCORE_BINDING_REQUIRED')
    receipt = read(out / 'score_receipt.json')
    if (receipt['task'] != TASK or receipt['freeze_sha256'] != completion['freeze_sha256']
            or receipt['review_sha256'] != completion['review_sha256']
            or receipt['predictions_sha256'] != digest(out / 'predictions.json')):
        raise ValueError('RETAINED_SCORE_RECEIPT_BINDING_MISMATCH')
    for relative, sha in completion['files'].items():
        if digest(out / relative) != sha:
            raise ValueError('RETAINED_SCORE_ARTIFACT_DRIFT')
    result = read(out / 'scores.json')
    print(canonical(result))
    return result


def evaluate(observed, predictions):
    expected = {(r['shot'], r['fraction']) for r in observed}
    if len(observed) != 96 or len(expected) != 96 or set(predictions) != set(ARMS):
        raise ValueError('ALL_ORIGINAL_PRED_SLOTS_AND_FOUR_ARMS_REQUIRED')
    supports = []
    for ps in predictions.values():
        if len(ps) != 96 or {(p['shot'], p['fraction']) for p in ps} != expected:
            raise ValueError('MISSING_PREDICTIONS_CANNOT_DROP_DENOMINATORS')
        supports.append({(p['shot'], p['fraction']): (p['supported'], p['reason']) for p in ps})
    if any(s != supports[0] for s in supports[1:]):
        raise ValueError('IDENTICAL_COMPARISON_SUPPORT_REQUIRED')
    panels, private = {}, {}
    for panel, conditions in PANELS.items():
        panels[panel], private[panel] = {}, {}
        for arm in ARMS:
            panels[panel][arm], private[panel][arm] = legacy.panel_metrics(observed, predictions[arm], conditions)
    primary = panels['primary']
    adequacy = {a: primary[a]['adequacy'] for a in ARMS}
    comparisons = {}
    for a, b in [('A1', 'E0'), ('A1', 'D0'), ('L1', 'D0'), ('L1', 'A1')]:
        gain = legacy.material_gain(primary[a], primary[b])
        gain['gain_status'] = gain['status']
        gain['candidate_adequacy'] = adequacy[a]
        gain['status'] = legacy.conjunction([gain['status'], adequacy[a]])
        comparisons[a+'_vs_'+b] = gain
    decision = decide(adequacy, comparisons, all(v['complete_scope'] for v in primary.values()))
    failures = {a: dict(Counter(p['prediction_status'] for p in ps if not p['numerical_qualified']))
                for a, ps in predictions.items()}
    extrapolation = {a: {panel: {'shots': len({p['shot'] for p in ps if p['condition'] in cs and p['feature_extrapolation']}),
                                 'features': sorted({f for p in ps if p['condition'] in cs for f in p['feature_extrapolation']})}
                          for panel, cs in PANELS.items()} for a, ps in predictions.items()}
    return {'task': TASK, **decision, 'adequacy_by_arm': adequacy, 'pairwise_complexity': comparisons,
            'panels': panels, 'prediction_failures_and_unsupported': failures, 'feature_extrapolation': extrapolation,
            'claims': list(md.CLAIMS), 'rights': source.RIGHTS, 'physical_validation': 'NOT_ESTABLISHED',
            'analytical_uncertainty': 'NOT_ESTABLISHED', 'scientific_score_passes': 1, 'outcome_joins': 1,
            'native_ewp_runs': 0, 'merge_authorized': False, 'production_adoption_authorized': False,
            'successor_authorized': False}, private


def main():
    parser = argparse.ArgumentParser(description=__doc__, allow_abbrev=False)
    sub = parser.add_subparsers(dest='stage', required=True)
    for name in ('prepare', 'develop', 'freeze', 'score', 'verify', 'report'):
        stage = sub.add_parser(name, allow_abbrev=False)
        stage.add_argument('--out', type=Path, required=True)
        if name == 'score': stage.add_argument('--review', type=Path, required=True)
    args = parser.parse_args()
    if args.stage in ('prepare', 'develop', 'freeze', 'score'):
        import resource
        soft, hard = resource.getrlimit(resource.RLIMIT_AS)
        cap = 12*1024**3
        resource.setrlimit(resource.RLIMIT_AS, (min(cap, soft) if soft > 0 else cap, hard))
    if args.stage == 'score': score(args.out, args.review)
    else: globals()[args.stage](args.out)


if __name__ == '__main__':
    main()
